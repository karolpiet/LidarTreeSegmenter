"""
LidarTreeSegmenter - OKNO
============

To samo, co PRZECIAGNIJ_TU_SKAN.bat, ale w okienku: widac etapy, pasek
postepu, czas, dziennik, podglad chmury na zywo (co program wlasnie
liczy) i obrazki drzew, w miare jak powstaja.

Kolejka na noc: wiele skanow liczonych po kolei (przycisk „Kolejka na
noc…” albo upuszczenie kilku plikow / folderu na LidarTreeSegmenter_PL.bat). Blad jednego
skanu nie zatrzymuje reszty, na koniec powstaje zestawienie wszystkich
powierzchni w Excelu (wyniki\\KOLEJKA_<data>.xlsx), a komputer moze sie
sam uspic albo wylaczyc.

Uruchamianie: kliknij dwa razy LidarTreeSegmenter_PL.bat (albo upusc na niego skan).
Bez dodatkowych bibliotek - okno jest zrobione w tkinter, ktory jest
w kazdej instalacji Pythona.
"""

import os
import queue
import re
import subprocess
import sys
import threading
import time
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

TU = Path(__file__).resolve().parent
ROOT = TU.parent                          # TU = LidarTreeSegmenter\\program
TST = ROOT / "tree_segmenter_tool"        # kopia w repozytorium
if not (TST / "02_detect_trees.py").exists():
    TST = ROOT.parent / "tree_segmenter_tool"   # stary uklad: folder obok
WYNIKI = ROOT / "wyniki"
sys.path.insert(0, str(TU))
import lts_uklad as U                    # noqa: E402
import lts_ustawienia as UST             # noqa: E402
import lts_jezyk as J                    # noqa: E402
T = J.t
BIBLIOTEKI = ["laspy", "lazrs", "numpy", "pandas", "scipy", "scikit-learn",
              "matplotlib", "openpyxl"]
IMPORTY = ("import laspy, lazrs, numpy, pandas, scipy, sklearn, matplotlib, "
           "openpyxl")

# --- kolory --------------------------------------------------------------
TLO = "#14181a"
PANEL = "#1d2326"
LINIA = "#2c3438"
TEKST = "#e6ebe8"
SZARY = "#8a9691"
AKCENT = "#7fb069"      # zielen lasu
DREWNO = "#c8893b"
BLAD = "#e05d5d"

# Etapy: (faza, numer KROK w dzienniku, opis, waga w pasku postepu)
ETAPY = [
    ("lib", None, T("Sprawdzenie bibliotek", "Checking libraries"), 1),
    ("drz", 1, T("Wczytanie skanu", "Loading the scan"), 4),
    ("drz", 2, T("Usuwanie szumu", "Noise removal"), 8),
    ("drz", 3, T("Model terenu", "Terrain model"), 4),
    ("drz", 4, T("Wysokość nad gruntem", "Height above ground"), 2),
    ("drz", 5, T("Wykrywanie pni", "Stem detection"), 22),
    ("drz", 6, T("Przypisanie punktów do drzew", "Assigning points to "
                 "trees"), 12),
    ("drz", 8, T("Pierśnica i wysokość", "DBH and height"), 6),
    ("drz", 9, T("Zapis tabeli drzew", "Saving the tree table"), 4),
    ("cz", 1, T("Wczytanie drzew", "Loading the trees"), 3),
    ("cz", 2, T("Śledzenie pni do korony", "Tracing stems to the crown"), 4),
    ("cz", 3, T("Cechy geometryczne", "Geometric features"), 12),
    ("cz", 4, T("Drewno czy liść", "Wood or leaf"), 5),
    ("cz", 5, T("Sprzątanie korony", "Crown clean-up"), 3),
    ("cz", 6, T("Zapis chmury części", "Saving the parts cloud"), 4),
    ("cz", "gpkg", T("GeoPackage do QGIS", "GeoPackage for QGIS"), 2),
    ("cz", "img", T("Obrazki wszystkich drzew", "Images of all trees"), 10),
    ("rap", "rap", T("Raport Excel i miąższość", "Excel report and volume"),
     2),
    ("rap", "3d", T("Widok 3D w przeglądarce", "3D view in the browser"), 3),
]
NAGLOWKI = {"lib": T("PRZYGOTOWANIE", "PREPARATION"),
            "drz": T("1  DRZEWA I PIERŚNICE", "1  TREES AND DBH"),
            "cz": T("2  PIEŃ / GAŁĘZIE / ULISTNIENIE",
                    "2  STEM / BRANCHES / FOLIAGE"),
            "rap": T("3  RAPORT I WIDOK 3D", "3  REPORT AND 3D VIEW")}
# wiersze dziennika z programów liczących - w obu językach (krok 1
# z tree_segmenter_tool pisze zawsze po polsku)
WZOR_KROK = re.compile(r"\[(?:KROK|STEP) (\d+)")
WZOR_RYS = re.compile(r"(?:narysowano|drawn) (\d+) (?:z|of) (\d+)")
WZOR_PROC = re.compile(r"^\s+(pien|galezie|ulistnienie|stem|branches|foliage)"
                       r"\s+([\d.]+)% (?:punktow|of points)")
WZOR_DRZEW = re.compile(r"(?:WYKRYTO DRZEW|TREES DETECTED):\s*(\d+)")
WZOR_V = re.compile(r"(?:MIAZSZOSC|VOLUME):\s*([\d.]+) m3/ha")
KOLOR_CZESCI = ["#4a5450", DREWNO, "#e0b070", AKCENT, "#8c94b3"]  # +podszyt


def interpreter():
    """python.exe obok pythonw.exe - zeby dzieci mialy normalne wyjscie."""
    exe = Path(sys.executable)
    if exe.name.lower() == "pythonw.exe" and (exe.parent / "python.exe").exists():
        return str(exe.parent / "python.exe")
    return str(exe)


def czuwaj(wlacz):
    """Windows: nie usypiaj komputera, dopóki ten wątek liczy."""
    if os.name != "nt":
        return
    try:
        import ctypes
        ES_CONTINUOUS, ES_SYSTEM_REQUIRED = 0x80000000, 0x00000001
        ctypes.windll.kernel32.SetThreadExecutionState(
            ES_CONTINUOUS | (ES_SYSTEM_REQUIRED if wlacz else 0))
    except Exception:                                # noqa: BLE001
        pass


def skany_w(sciezki):
    """Pliki .laz/.las z podanych plików i folderów (z podfolderami),
    bez folderu wyników LidarTreeSegmenter."""
    out = []
    for s in sciezki:
        p = Path(str(s).strip().strip('"'))
        if p.is_file() and p.suffix.lower() in (".laz", ".las"):
            out.append(p)
        elif p.is_dir():
            for f in sorted(p.rglob("*")):
                if f.suffix.lower() in (".laz", ".las") and f.is_file() \
                        and WYNIKI not in f.parents:
                    out.append(f)
    widziane, wynik = set(), []
    for f in out:
        k = str(f.resolve()).lower()
        if k not in widziane:
            widziane.add(k)
            wynik.append(f)
    return wynik


def podsumowanie(folder):
    """Wartości z 01_RAPORT/00_PODSUMOWANIE.txt (albo 00_SUMMARY.txt
    z wersji angielskiej): {nazwa: wartość jako liczba albo tekst}.
    Liczby są od razu zamieniane (1 214,5 po polsku, 1,214.5 po ang.)."""
    p = U.znajdz(folder, "podsumowanie")
    try:
        txt = p.read_text(encoding="utf-8")
    except (OSError, AttributeError):
        return {}
    en = p.name.lower() == "00_summary.txt"
    d = {}
    for ln in txt.splitlines():
        m = re.match(r"^  (\S.*?)\s*: (.+)$", ln)
        if m and m.group(1) not in d:
            d[m.group(1)] = liczba(m.group(2).strip(), en)
    return d


def liczba(tekst, en=False):
    """'1 214 m² (0,1215 ha)' -> 1214.0 ('1,214 m²' gdy en); brak -> None."""
    if en:
        m = re.match(r"^\s*(-?[\d,]+(?:\.\d+)?)", tekst or "")
        tekst = m.group(1).replace(",", "").replace(".", ",") if m else ""
    m = re.match(r"^\s*(-?[\d\s\u00a0]+(?:,\d+)?)", tekst or "")
    if not m:
        return None
    try:
        return float(re.sub(r"[\s\u00a0]", "", m.group(1)).replace(",", "."))
    except ValueError:
        return None


KROTKO_KLASA = {"auto": T("automatycznie", "automatic"),
                "młody": T("młody", "young"),
                "dojrzały": T("dojrzały", "mature")}
KROTKO_GAT = {"auto": T("rozpoznaj", "detect"),
              "świerk": T("same świerki", "all spruce"),
              "buk": T("same buki", "all beech")}
PO_KOLEJCE = {T("nic nie rób", "do nothing"): None,
              T("uśpij komputer", "sleep"): "usp",
              T("wyłącz komputer", "shut down"): "wyl"}
# stan pozycji w kolejce: wewnętrznie po polsku, pokazywany w języku okna
STAN = {"czeka": T("czeka", "waiting"), "liczy się": T("liczy się",
        "running"), "gotowe": T("gotowe", "done"), "błąd": T("błąd", "error"),
        "zatrzymano": T("zatrzymano", "stopped")}
# wiersze 00_PODSUMOWANIE.txt do zestawienia kolejki: (po polsku, po ang.)
KLUCZE_PODS = [
    ("Powierzchnia próbna", "Sample plot area"),
    ("Liczba drzew", "Number of trees"),
    ("Liczba drzew na 1 ha (N)", "Trees per hectare (N)"),
    ("Pierśnica przeciętna (Dg", "Quadratic mean diameter (Dg)"),
    ("Wysokość Loreya", "Lorey's height"),
    ("Pierśnicowe pole przekroju (G)", "Basal area (G)"),
    ("Miąższość grubizny – powierzchnia", "Merchantable volume ≥7 cm – plot"),
    ("Miąższość grubizny (V)", "Merchantable volume ≥7 cm (V)"),
    ("Pominięte jako krzewy", "Excluded as shrubs"),
    ("Złomy", "Snags")]


def wartosc_pods(ps, klucz):
    """Pierwsza wartość z podsumowania, której nazwa zaczyna się od klucza
    (polskiego albo angielskiego)."""
    return next((v for n, v in ps.items() if n.startswith(klucz)), None)


PUSTE_OBRAZKI = T("Tu pojawią się obrazki drzew,\nw miarę jak będą powstawać.",
                  "Tree images will appear here\nas they are created.")
ZATRZYMANO = T("Zatrzymano.", "Stopped.")
BLAD_LIB = T("Instalacja bibliotek się nie udała.",
             "Library installation failed.")


class Okno:
    def __init__(self, root, skan=None):
        self.root = root
        self.q = queue.Queue()
        self.proc = None
        self.przerwane = False
        self.t_start = None
        self.dziennik_caly = []
        self.stan = ["czeka"] * len(ETAPY)
        self.ulamek = [0.0] * len(ETAPY)
        self.biezacy = None
        self.obrazek = None
        self.ostatni_obrazek = None
        self.wynik_folder = None
        self.plik_zywo = None
        self.t_zywo = 0.0
        self.zywo = None              # (figura, płótno, osie) po 1. użyciu
        self.kolejka = []             # [{skan, klasa, gatunek, stan, ...}]
        self.w_kolejce = False
        self.okno_kolejki = None
        self.ost_v = None
        self.ost_blad = ""

        root.title(J.NAZWA)
        root.configure(bg=TLO)
        root.geometry(f"1180x{min(900, root.winfo_screenheight() - 60)}")
        if os.name == "nt" and root.winfo_screenheight() < 900:
            root.state("zoomed")              # maly ekran - na caly ekran
        root.minsize(980, 640)
        self._style()
        self._zbuduj()
        skany = skany_w(skan) if skan else []
        if len(skany) == 1 and Path(str(skan[0]).strip('"')).is_file():
            self.var_skan.set(str(skany[0]))
            self._pokaz_ustawienia(skany[0].stem)
            self.root.after(400, self.start)
        elif skany:                      # kilka plików albo folder
            self.root.after(300, lambda: self.pokaz_kolejke(skany))
        self.root.after(100, self._petla)

    # ------------------------------------------------------------------
    def _style(self):
        s = ttk.Style()
        s.theme_use("clam")
        s.configure(".", background=TLO, foreground=TEKST,
                    font=("Segoe UI", 10))
        s.configure("Panel.TFrame", background=PANEL)
        s.configure("TLabel", background=TLO, foreground=TEKST)
        s.configure("Panel.TLabel", background=PANEL, foreground=TEKST)
        s.configure("Szary.TLabel", background=PANEL, foreground=SZARY,
                    font=("Segoe UI", 9))
        s.configure("Tytul.TLabel", background=TLO, foreground=TEKST,
                    font=("Segoe UI Semibold", 22))
        s.configure("Pod.TLabel", background=TLO, foreground=SZARY,
                    font=("Segoe UI", 10))
        s.configure("Duzy.TLabel", background=PANEL, foreground=TEKST,
                    font=("Segoe UI Semibold", 26))
        s.configure("Proc.TLabel", background=TLO, foreground=AKCENT,
                    font=("Segoe UI Semibold", 13))
        s.configure("Akcent.TButton", background=AKCENT, foreground="#0e1210",
                    font=("Segoe UI Semibold", 11), padding=(18, 8),
                    borderwidth=0)
        s.map("Akcent.TButton", background=[("disabled", LINIA),
                                            ("active", "#94c47f")],
              foreground=[("disabled", SZARY)])
        s.configure("TButton", background=PANEL, foreground=TEKST,
                    padding=(14, 7), borderwidth=1, bordercolor=LINIA)
        s.map("TButton", background=[("active", LINIA)],
              foreground=[("disabled", "#55605b")])
        s.configure("TEntry", fieldbackground=PANEL, foreground=TEKST,
                    bordercolor=LINIA, insertcolor=TEKST, padding=6)
        s.configure("TCombobox", fieldbackground=PANEL, background=PANEL,
                    foreground=TEKST, arrowcolor=TEKST, bordercolor=LINIA,
                    padding=4)
        s.map("TCombobox", fieldbackground=[("readonly", PANEL)],
              foreground=[("readonly", TEKST)],
              selectbackground=[("readonly", PANEL)],
              selectforeground=[("readonly", TEKST)])
        self.root.option_add("*TCombobox*Listbox.background", PANEL)
        self.root.option_add("*TCombobox*Listbox.foreground", TEKST)
        s.configure("Ciemny.TNotebook", background=PANEL, borderwidth=0)
        s.configure("Ciemny.TNotebook.Tab", background=TLO, foreground=SZARY,
                    padding=(10, 4), borderwidth=0)
        s.map("Ciemny.TNotebook.Tab", background=[("selected", LINIA)],
              foreground=[("selected", TEKST)])
        s.configure("Zielony.Horizontal.TProgressbar", troughcolor=PANEL,
                    background=AKCENT, bordercolor=PANEL, lightcolor=AKCENT,
                    darkcolor=AKCENT, thickness=14)

    def _zbuduj(self):
        g = ttk.Frame(self.root, padding=(22, 16, 22, 6))
        g.pack(fill="x")
        ttk.Label(g, text=J.NAZWA, style="Tytul.TLabel").pack(side="left")
        ttk.Label(g, text=T("   drzewa  ·  pierśnice  ·  pień / gałęzie / "
                            "ulistnienie", "   trees  ·  DBH  ·  stem / "
                            "branches / foliage"),
                  style="Pod.TLabel").pack(side="left", pady=(10, 0))

        # --- wybor skanu ---
        w = ttk.Frame(self.root, padding=(22, 6))
        w.pack(fill="x")
        self.var_skan = tk.StringVar()
        ttk.Entry(w, textvariable=self.var_skan).pack(side="left", fill="x",
                                                      expand=True)
        ttk.Button(w, text=T("Wybierz skan…", "Choose scan…"),
                   command=self.wybierz).pack(
            side="left", padx=(8, 0))
        self.b_start = ttk.Button(w, text="▶  Start", style="Akcent.TButton",
                                  command=self.start)
        self.b_start.pack(side="left", padx=(8, 0))
        self.b_stop = ttk.Button(w, text=T("■  Zatrzymaj", "■  Stop"),
                                 command=self.stop,
                                 state="disabled")
        self.b_stop.pack(side="left", padx=(8, 0))
        self.b_kolejka = ttk.Button(w, text=T("📋  Kolejka na noc…",
                                              "📋  Overnight queue…"),
                                    command=self.pokaz_kolejke)
        self.b_kolejka.pack(side="left", padx=(8, 0))

        # --- klasa drzewostanu i gatunek ---
        o = ttk.Frame(self.root, padding=(22, 2))
        o.pack(fill="x")
        self.KLASY = {T("automatycznie (z pierśnic)",
                        "automatic (from DBH)"): "auto",
                      T("młody (pierśnice ok. 8–20 cm)",
                        "young (DBH approx. 8–20 cm)"): "młody",
                      T("dojrzały (pierśnice ponad 20 cm)",
                        "mature (DBH above 20 cm)"): "dojrzały"}
        self.GATUNKI = {T("rozpoznaj z kształtu (wstępnie)",
                          "detect from shape (preliminary)"): "auto",
                        T("same świerki", "all spruce"): "świerk",
                        T("same buki", "all beech"): "buk"}
        ttk.Label(o, text=T("Drzewostan:", "Stand:"),
                  style="Pod.TLabel").pack(side="left")
        self.var_klasa = tk.StringVar(value=next(iter(self.KLASY)))
        ttk.Combobox(o, textvariable=self.var_klasa, values=list(self.KLASY),
                     state="readonly", width=30).pack(side="left",
                                                      padx=(6, 18))
        ttk.Label(o, text=T("Gatunek:", "Species:"),
                  style="Pod.TLabel").pack(side="left")
        self.var_gatunek = tk.StringVar(value=next(iter(self.GATUNKI)))
        ttk.Combobox(o, textvariable=self.var_gatunek,
                     values=list(self.GATUNKI), state="readonly",
                     width=30).pack(side="left", padx=(6, 0))
        self.l_kolejka = ttk.Label(o, text="", style="Proc.TLabel")
        self.l_kolejka.pack(side="right")

        # --- ścieżka skanera (trajektoria) - puste pole: szukaj samemu ---
        sc = ttk.Frame(self.root, padding=(22, 4, 22, 0))
        sc.pack(fill="x")
        ttk.Label(sc, text=T("Ścieżka skanera:", "Scanner path:"),
                  style="Pod.TLabel").pack(side="left")
        self.var_sciezka = tk.StringVar()
        ttk.Entry(sc, textvariable=self.var_sciezka, width=48).pack(
            side="left", padx=(6, 0))
        ttk.Button(sc, text=T("Wybierz…", "Choose…"),
                   command=self.wybierz_sciezke).pack(side="left",
                                                      padx=(8, 0))
        ttk.Button(sc, text=T("Szukaj sam", "Find automatically"),
                   command=self.sciezka_auto).pack(side="left", padx=(8, 0))
        self.l_sciezka = ttk.Label(sc, text="", style="Szary.TLabel")
        self.l_sciezka.pack(side="left", padx=(10, 0))

        # --- pasek postepu ---
        p = ttk.Frame(self.root, padding=(22, 10, 22, 4))
        p.pack(fill="x")
        self.pasek = ttk.Progressbar(p, style="Zielony.Horizontal.TProgressbar",
                                     maximum=1000)
        self.pasek.pack(side="left", fill="x", expand=True)
        self.l_proc = ttk.Label(p, text="0%", style="Proc.TLabel", width=5,
                                anchor="e")
        self.l_proc.pack(side="left", padx=(10, 0))
        self.l_czas = ttk.Label(p, text="00:00", style="Pod.TLabel", width=7,
                                anchor="e")
        self.l_czas.pack(side="left")
        self.l_status = ttk.Label(self.root, text=T(
            "Wybierz skan i kliknij Start.", "Choose a scan and click Start."),
                                  style="Pod.TLabel", padding=(22, 0, 22, 8))
        self.l_status.pack(fill="x")

        # --- przyciski na dole ---
        dol = ttk.Frame(self.root, padding=(22, 0, 22, 16))
        dol.pack(side="bottom", fill="x")
        self.b_folder = ttk.Button(dol, text=T("📂  Folder wyników",
                                               "📂  Results folder"),
                                   command=self.otworz_folder, state="disabled")
        self.b_folder.pack(side="left")
        self.b_podglad = ttk.Button(dol, text=T("🌲  Obrazki drzew",
                                                "🌲  Tree images"),
                                    command=self.otworz_podglad,
                                    state="disabled")
        self.b_podglad.pack(side="left", padx=(8, 0))
        self.b_excel = ttk.Button(dol, text=T("📊  Raport Excel",
                                              "📊  Excel report"),
                                  command=self.otworz_excel, state="disabled")
        self.b_excel.pack(side="left", padx=(8, 0))
        self.b_3d = ttk.Button(dol, text=T("🧊  Widok 3D", "🧊  3D view"),
                               command=self.otworz_3d, state="disabled")
        self.b_3d.pack(side="left", padx=(8, 0))
        self.b_spr = ttk.Button(dol, text=T("✅  Sprawdź drzewa",
                                            "✅  Review trees"),
                                command=self.sprawdz_drzewa, state="disabled")
        self.b_spr.pack(side="left", padx=(8, 0))
        self.b_przelicz = ttk.Button(dol, text=T("🔄  Przelicz raport",
                                                 "🔄  Recalculate report"),
                                     command=self.przelicz_raport,
                                     state="disabled")
        self.b_przelicz.pack(side="left", padx=(8, 0))
        ttk.Button(dol, text=T("Kopiuj dziennik", "Copy log"),
                   command=self.kopiuj).pack(
            side="right")

        # --- dziennik ---
        d = ttk.Frame(self.root, padding=(22, 0, 22, 8))
        d.pack(side="bottom", fill="x")
        self.log = tk.Text(d, height=6, bg="#0f1214", fg="#b9c4bf",
                           insertbackground=TEKST, relief="flat",
                           font=("Consolas", 9), wrap="none",
                           highlightthickness=1, highlightbackground=LINIA)
        sb = ttk.Scrollbar(d, command=self.log.yview)
        self.log.configure(yscrollcommand=sb.set, state="disabled")
        self.log.pack(side="left", fill="both", expand=True)
        sb.pack(side="left", fill="y")
        self.log.tag_configure("blad", foreground=BLAD)
        self.log.tag_configure("info", foreground=AKCENT)

        # --- srodek: etapy | podglad ---
        sr = ttk.Frame(self.root, padding=(22, 0, 22, 8))
        sr.pack(fill="both", expand=True)
        sr.columnconfigure(0, weight=0, minsize=470)
        sr.columnconfigure(1, weight=1)
        sr.rowconfigure(0, weight=1)

        lewy = ttk.Frame(sr, style="Panel.TFrame", padding=(16, 12))
        lewy.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        self.wiersze = []
        faza_pop = None
        # dwie kolumny etapów, żeby lista mieściła się na małym ekranie
        kolumny = [ttk.Frame(lewy, style="Panel.TFrame") for _ in range(2)]
        kolumny[0].pack(side="left", fill="both", expand=True, anchor="n")
        kolumny[1].pack(side="left", fill="both", expand=True, anchor="n",
                        padx=(12, 0))
        for i, (faza, _, opis, _) in enumerate(ETAPY):
            kol = kolumny[0 if faza in ("lib", "drz") else 1]
            if faza != faza_pop:
                ttk.Label(kol, text=NAGLOWKI[faza], style="Szary.TLabel").pack(
                    anchor="w", pady=(8 if faza in ("drz", "rap") else 0, 2))
                faza_pop = faza
            r = tk.Frame(kol, bg=PANEL)
            r.pack(fill="x", pady=0)
            ikona = tk.Label(r, text="○", bg=PANEL, fg=SZARY, width=2,
                             font=("Segoe UI Symbol", 11))
            ikona.pack(side="left")
            napis = tk.Label(r, text=opis, bg=PANEL, fg=SZARY, anchor="w",
                             font=("Segoe UI", 9))
            napis.pack(side="left", fill="x", expand=True)
            dod = tk.Label(r, text="", bg=PANEL, fg=SZARY,
                           font=("Segoe UI", 9))
            dod.pack(side="right")
            self.wiersze.append((ikona, napis, dod))

        prawy = ttk.Frame(sr, style="Panel.TFrame", padding=(16, 12))
        prawy.grid(row=0, column=1, sticky="nsew")
        prawy.rowconfigure(1, weight=1)
        prawy.columnconfigure(0, weight=1)
        gora = ttk.Frame(prawy, style="Panel.TFrame")
        gora.grid(row=0, column=0, sticky="ew")
        self.kafle = {}
        for k, (klucz, nazwa, kolor) in enumerate((
                ("drzew", T("drzew", "trees"), TEKST),
                ("pień", T("pień", "stem"), DREWNO),
                ("gałęzie", T("gałęzie", "branches"), "#e0b070"),
                ("ulistnienie", T("ulistnienie", "foliage"), AKCENT),
                ("m³/ha", "m³/ha", "#9cc3e6"))):
            f = tk.Frame(gora, bg=PANEL)
            f.pack(side="left", expand=True, fill="x")
            v = tk.Label(f, text="–", bg=PANEL, fg=kolor,
                         font=("Segoe UI Semibold", 24))
            v.pack()
            tk.Label(f, text=nazwa, bg=PANEL, fg=SZARY,
                     font=("Segoe UI", 9)).pack()
            self.kafle[klucz] = v
        self.zakladki = ttk.Notebook(prawy, style="Ciemny.TNotebook")
        self.zakladki.grid(row=1, column=0, sticky="nsew", pady=(10, 0))
        self.f_zywo = tk.Frame(self.zakladki, bg=PANEL)
        f_obr = tk.Frame(self.zakladki, bg=PANEL)
        self.zakladki.add(self.f_zywo, text=T("  Podgląd na żywo  ",
                                              "  Live preview  "))
        self.zakladki.add(f_obr, text=T("  Obrazki drzew  ",
                                        "  Tree images  "))
        self.l_zywo_pusty = tk.Label(
            self.f_zywo, bg=PANEL, fg=SZARY, font=("Segoe UI", 10),
            text=T("Tu zobaczysz chmurę punktów w trakcie liczenia:\n"
                   "szum, grunt, znalezione pnie, drzewa, pień / gałęzie / "
                   "liście.", "Here you will see the point cloud while it is "
                   "being processed:\nnoise, ground, stems found, trees, "
                   "stem / branches / leaves."))
        self.l_zywo_pusty.pack(fill="both", expand=True)
        self.platno = tk.Label(f_obr, bg=PANEL, fg=SZARY,
                               text=PUSTE_OBRAZKI,
                               font=("Segoe UI", 10))
        self.platno.pack(fill="both", expand=True)
        self.l_obrazek = ttk.Label(prawy, text="", style="Szary.TLabel")
        self.l_obrazek.grid(row=2, column=0, sticky="w")

    # ------------------------------------------------------------------
    def _pokaz_ustawienia(self, nazwa):
        """Wybór z poprzedniego liczenia tego skanu."""
        u = UST.wczytaj(WYNIKI / nazwa, nazwa)
        for zm, slownik, klucz in ((self.var_klasa, self.KLASY, "klasa"),
                                   (self.var_gatunek, self.GATUNKI,
                                    "gatunek")):
            for opis, w in slownik.items():
                if w == u.get(klucz):
                    zm.set(opis)
        self.var_sciezka.set(u.get("sciezka", ""))
        self._opisz_sciezke()
        self._przyciski(WYNIKI / nazwa, nazwa)

    def _opisz_sciezke(self):
        """Pod polem: jaki plik ścieżki skanera zostanie użyty."""
        wyb = self.var_sciezka.get().strip().strip('"')
        if wyb:
            self.l_sciezka.config(text=T("wybrana ręcznie", "chosen by hand")
                                  if Path(wyb).is_file() else
                                  T("nie ma takiego pliku!",
                                    "no such file!"))
            return
        skan = self.var_skan.get().strip().strip('"')
        if not skan:
            self.l_sciezka.config(text="")
            return
        try:
            import lts_sciezka as SC
            nazwa = Path(skan).stem
            kat = WYNIKI / nazwa
            # samo szukanie, bez ręcznego wyboru zapisanego wcześniej
            p = SC.znajdz(kat / "_brak_", nazwa, plik_skanu=skan) or \
                SC.znajdz(kat, nazwa, plik_skanu=skan)
        except Exception:                            # noqa: BLE001
            p = None
        self.l_sciezka.config(text=T(f"znaleziona sama: {p.name}",
                                     f"found automatically: {p.name}")
                              if p else T("nie znaleziono (bez niej też "
                                          "policzy)", "not found (works "
                                          "without it too)"))

    def wybierz_sciezke(self):
        skan = self.var_skan.get().strip().strip('"')
        p = filedialog.askopenfilename(
            title=T("Wybierz plik ścieżki skanera",
                    "Choose the scanner path file"),
            filetypes=[(T("Ścieżka skanera", "Scanner path"),
                        "*.las *.laz *.txt *.csv *.xyz *.pts *.asc *.ply"),
                       (T("Wszystkie pliki", "All files"), "*.*")],
            initialdir=str(Path(skan).parent) if skan else
            str(ROOT.parent))
        if p:
            self.var_sciezka.set(str(Path(p)))
            self._opisz_sciezke()

    def sciezka_auto(self):
        self.var_sciezka.set("")
        self._opisz_sciezke()

    def _przyciski(self, folder, nazwa):
        """Włącz przyciski, do których są już wyniki tego skanu."""
        if self.w_kolejce or str(self.b_stop["state"]) == "normal":
            return
        ist = folder.exists()
        self.wynik_folder, self.nazwa = (folder, nazwa) if ist else \
            (self.wynik_folder, getattr(self, "nazwa", nazwa))
        self.b_folder.config(state="normal" if ist else "disabled")
        self.b_podglad.config(state="normal" if (folder / U.OBRAZY).exists()
                              else "disabled")
        self.b_excel.config(state="normal" if U.znajdz(folder, "raport",
                                                       nazwa) else "disabled")
        html = U.znajdz(folder, "widok", nazwa) is not None
        self.b_3d.config(state="normal" if html else "disabled")
        self.b_spr.config(state="normal" if html else "disabled")
        self.b_przelicz.config(state="normal" if (
            folder / U.CHMURY / f"{nazwa}_czesci.laz").exists()
            else "disabled")

    def wybierz(self):
        p = filedialog.askopenfilename(
            title=T("Wybierz skan", "Choose a scan"),
            filetypes=[(T("Chmury punktów", "Point clouds"), "*.laz *.las"),
                       (T("Wszystkie pliki", "All files"), "*.*")],
            initialdir=str(ROOT.parent))
        if p:
            self.var_skan.set(p)
            self.var_sciezka.set("")
            self._pokaz_ustawienia(Path(p).stem)

    def _ustaw_skan(self, skan):
        self.skan = Path(skan)
        self.nazwa = self.skan.stem
        self.wynik_folder = WYNIKI / self.nazwa
        self.plik_zywo = self.wynik_folder / U.DZIENNIK / "na_zywo.npz"

    def _zapisz_ustawienia(self, klasa, gatunek, sciezka=None):
        """Zapisuje wybór dla skanu; True, gdy różni się od poprzedniego.
        sciezka: plik ścieżki skanera ("" = szukaj samemu, None = bez
        zmian)."""
        stare = UST.wczytaj(self.wynik_folder, self.nazwa)
        UST.zapisz(self.wynik_folder, self.nazwa, klasa=klasa,
                   gatunek=gatunek, sciezka=sciezka)
        return (stare.get("klasa", "auto") != klasa
                or stare.get("gatunek", "auto") != gatunek
                or (sciezka is not None
                    and stare.get("sciezka", "") != sciezka))

    def start(self, tryb="wszystko"):
        skan = Path(self.var_skan.get().strip().strip('"'))
        if not skan.is_file():
            messagebox.showerror(J.NAZWA, T(f"Nie ma takiego pliku:\n{skan}",
                                            f"No such file:\n{skan}"))
            return
        self._ustaw_skan(skan)
        self._zapisz_ustawienia(self.KLASY[self.var_klasa.get()],
                                self.GATUNKI[self.var_gatunek.get()],
                                self.var_sciezka.get().strip().strip('"'))
        self._wyczysc()
        self.b_start.config(state="disabled")
        self.b_kolejka.config(state="disabled")
        self.b_stop.config(state="normal")
        self.b_folder.config(state="normal" if self.wynik_folder.exists()
                             else "disabled")
        self.t_start = time.time()
        threading.Thread(target=self._praca, args=(tryb,),
                         daemon=True).start()

    def przelicz_raport(self):
        """Tylko raport i widok 3D - np. po sprawdzeniu drzew albo
        zmianie gatunku. Bez liczenia chmury od nowa."""
        self.start("raport")

    def _wyczysc(self):
        """Okno gotowe na nowy skan."""
        self.t_zywo = 0.0
        self.stan = ["czeka"] * len(ETAPY)
        self.ulamek = [0.0] * len(ETAPY)
        self.biezacy = None
        self.przerwane = False
        self.dziennik_caly = []
        self.ostatni_obrazek = None
        for v in self.kafle.values():
            v.config(text="–")
        self.platno.config(image="", text=PUSTE_OBRAZKI)
        self.log.config(state="normal"); self.log.delete("1.0", "end")
        self.log.config(state="disabled")
        for w_ in self.wiersze:
            w_[2].config(text="")
        self._rysuj_etapy()

    def stop(self):
        if self.w_kolejce and not messagebox.askyesno(
                J.NAZWA, T("Zatrzymać całą kolejkę?\n(Policzone skany "
                           "zostaną, zestawienie się zapisze.)", "Stop the "
                           "whole queue?\n(Finished scans are kept and the "
                           "summary is saved.)")):
            return
        self.przerwane = True
        if self.proc and self.proc.poll() is None:
            self.proc.terminate()
        self.q.put(("status", ZATRZYMANO))

    def otworz_folder(self):
        if self.wynik_folder and self.wynik_folder.exists():
            os.startfile(self.wynik_folder) if os.name == "nt" else None

    def otworz_podglad(self):
        p = self.wynik_folder / U.OBRAZY if self.wynik_folder else None
        if p and p.exists() and os.name == "nt":
            os.startfile(p)

    def otworz_excel(self):
        if not self.wynik_folder:
            return
        p = U.znajdz(self.wynik_folder, "raport", self.nazwa)
        if p and os.name == "nt":
            os.startfile(p)

    def otworz_3d(self):
        if not self.wynik_folder:
            return
        p = U.znajdz(self.wynik_folder, "widok", self.nazwa)
        if p and os.name == "nt":
            os.startfile(p)

    def sprawdz_drzewa(self):
        """Widok 3D od razu w trybie sprawdzania drzew po kolei."""
        if not self.wynik_folder:
            return
        p = U.znajdz(self.wynik_folder, "sprawdzanie", self.nazwa) or \
            U.znajdz(self.wynik_folder, "widok", self.nazwa)
        if p and os.name == "nt":
            os.startfile(p)
        messagebox.showinfo(J.NAZWA, T(
            "Sprawdzanie drzew otworzy się w przeglądarce.\n\n"
            "Klawisze: 1 drzewo OK, 2 krzak, 3 złom, 4 świerk, 5 buk,\n"
            "→ pomiń, ← wstecz.\n\nNa koniec kliknij „Zapisz sprawdzenie” "
            f"i zapisz plik w folderze\n{self.wynik_folder / U.RAPORT}\n"
            "Potem kliknij tu „Przelicz raport”.",
            "The tree review opens in the browser.\n\n"
            "Keys: 1 tree OK, 2 shrub, 3 snag, 4 spruce, 5 beech,\n"
            "→ skip, ← back.\n\nAt the end click “Save review” "
            f"and save the file in the folder\n{self.wynik_folder / U.RAPORT}"
            "\nThen click “Recalculate report” here."))

    def kopiuj(self):
        self.root.clipboard_clear()
        self.root.clipboard_append("".join(self.dziennik_caly))
        self.l_status.config(text=T("Dziennik skopiowany - wklej go "
                                    "Claude'owi.", "Log copied - paste it "
                                    "to Claude."))

    # ------------------------------------------------------------------
    # Praca w osobnym watku: kolejne programy, wyjscie linia po linii.
    def _uruchom(self, faza, polecenie):
        env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1")
        if self.plik_zywo:
            env["LTS_LIVE"] = str(self.plik_zywo)
        flagi = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
        self.proc = subprocess.Popen(
            polecenie, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            cwd=str(TU), env=env, creationflags=flagi, bufsize=1,
            text=True, encoding="utf-8", errors="replace")
        for linia in self.proc.stdout:
            self.q.put(("linia", faza, linia))
            m = WZOR_V.search(linia)
            if m:
                self.ost_v = float(m.group(1))
            if re.search(r"Error|ERROR|BLAD|BŁĄD", linia):
                self.ost_blad = linia.strip()[:200]
        return self.proc.wait()

    def _biblioteki(self, py):
        self.q.put(("etap", 0))
        if subprocess.run([py, "-c", IMPORTY], capture_output=True,
                          creationflags=subprocess.CREATE_NO_WINDOW
                          if os.name == "nt" else 0).returncode != 0:
            self.q.put(("info", T("Brakuje bibliotek - instaluję "
                                  "(tylko za pierwszym razem)…\n",
                                  "Libraries are missing - installing them "
                                  "(first time only)…\n")))
            if self._uruchom("lib", [py, "-m", "pip", "install"] +
                             BIBLIOTEKI) != 0:
                return False
        self.q.put(("gotowe", 0))
        return True

    def _praca(self, tryb="wszystko"):
        py = interpreter()
        czuwaj(True)
        try:
            if not self._biblioteki(py):
                return self.q.put(("koniec", False, BLAD_LIB))
            ok, tekst = self._licz_skan(py, tryb)
            self.q.put(("koniec", ok, tekst))
        except Exception as e:                       # noqa: BLE001
            self.q.put(("koniec", False, T(f"Błąd: {e}", f"Error: {e}")))
        finally:
            czuwaj(False)

    @staticmethod
    def _nowszy(plik, *od):
        """plik istnieje i jest nowszy od wszystkich istniejących z od."""
        if not plik or not plik.exists():
            return False
        t = plik.stat().st_mtime
        return all(t >= o.stat().st_mtime for o in od if o and o.exists())

    def _licz_skan(self, py, tryb="wszystko", pomijaj=False):
        """Liczy jeden skan (self.skan). Zwraca (udało się, opis).
        tryb "raport": tylko raport i widok 3D. pomijaj: etapy, których
        wyniki są już nowsze od danych wejściowych, nie liczą się od nowa."""
        U.uporzadkuj(self.wynik_folder)
        seg = U.plik(self.wynik_folder, f"{self.nazwa}_segmenty.laz")
        czesci = U.plik(self.wynik_folder, f"{self.nazwa}_czesci.laz")
        if tryb == "raport":
            if not czesci.exists():
                return False, T("Nie ma jeszcze wyników etapu 2 - kliknij "
                                "Start, żeby policzyć cały skan.",
                                "There are no stage 2 results yet - click "
                                "Start to process the whole scan.")
            self.q.put(("pomin", "drz"))
            self.q.put(("pomin", "cz"))
        else:
            # 1. drzewa (tree_segmenter_tool) - albo wynik z wczesniej
            stare = TST / "results" / self.nazwa
            if not seg.exists() and (stare / seg.name).exists():
                self.q.put(("info", T("Przenoszę wcześniejszy wynik z "
                                      "tree_segmenter_tool\\results…\n",
                                      "Moving the earlier result from "
                                      "tree_segmenter_tool\\results…\n")))
                self.wynik_folder.mkdir(parents=True, exist_ok=True)
                for f in stare.iterdir():
                    if f.is_file():
                        f.replace(self.wynik_folder / f.name)
                U.uporzadkuj(self.wynik_folder)
                seg = U.plik(self.wynik_folder, seg.name)
            if seg.exists():
                self.q.put(("info", T("Drzewa były już znalezione wcześniej "
                                      "- pomijam etap 1.\n", "Trees were "
                                      "already found earlier - skipping "
                                      "stage 1.\n")))
                self.q.put(("pomin", "drz"))
            else:
                kod = self._uruchom("drz", [py, "-u",
                                            str(TU / "05_wykryj_drzewa.py"),
                                            str(self.skan), "--wyniki",
                                            str(WYNIKI)])
                if self.przerwane:
                    return False, ZATRZYMANO
                U.uporzadkuj(self.wynik_folder)
                seg = U.plik(self.wynik_folder, seg.name)
                if kod != 0 or not seg.exists():
                    return False, T("Etap 1 się nie udał.",
                                    "Stage 1 failed.")
                self.q.put(("faza_koniec", "drz"))

            # 2. pien / galezie / ulistnienie
            czesci = U.plik(self.wynik_folder, czesci.name)
            if pomijaj and self._nowszy(czesci, seg):
                self.q.put(("info", T("Pień / gałęzie / liście były już "
                                      "policzone - pomijam etap 2.\n",
                                      "Stem / branches / leaves were already "
                                      "computed - skipping stage 2.\n")))
                self.q.put(("pomin", "cz"))
            else:
                kod = self._uruchom("cz", [py, "-u", str(
                    TU / "10_pien_galezie_liscie.py"), str(seg), "--wyniki",
                    str(WYNIKI)])
                if self.przerwane:
                    return False, ZATRZYMANO
                if kod != 0:
                    return False, T("Etap 2 się nie udał.",
                                    "Stage 2 failed.")
                self.q.put(("faza_koniec", "cz"))
                U.uporzadkuj(self.wynik_folder)
                czesci = U.plik(self.wynik_folder, czesci.name)

        # 3. raport: Excel + 00_PODSUMOWANIE.txt z miąższością
        rap = self.wynik_folder / U.RAPORT
        xlsx = U.pliki(self.wynik_folder, "raport", self.nazwa)
        wej = [czesci] + [p for p in (rap.glob(f"{self.nazwa}_*") if
                                      rap.exists() else [])
                          if p.suffix.lower() in (".csv", ".gpx", ".xlsx")
                          and "_RAPORT" not in p.name
                          and "_REPORT" not in p.name]
        if pomijaj and xlsx and self._nowszy(xlsx[-1], *wej):
            self.q.put(("info", T("Raport był już zrobiony - pomijam.\n",
                                  "The report was already made - "
                                  "skipping.\n")))
        else:
            self.q.put(("etap", self._indeks("rap", "rap")))
            kod = self._uruchom("rap", [py, "-u", str(TU / "30_raport.py"),
                                        str(self.wynik_folder)])
            if self.przerwane:
                return False, ZATRZYMANO
            if kod != 0:
                return False, T("Raport się nie udał.", "The report failed.")
        # 4. widok 3D (HTML) - brak widoku nie psuje reszty wyników
        html = U.znajdz(self.wynik_folder, "widok", self.nazwa)
        raport_csv = U.plik(self.wynik_folder,
                            f"{self.nazwa}_drzewa_raport.csv")
        if pomijaj and self._nowszy(html, czesci, raport_csv):
            self.q.put(("info", T("Widok 3D był już zrobiony - pomijam.\n",
                                  "The 3D view was already made - "
                                  "skipping.\n")))
        else:
            self.q.put(("etap", self._indeks("rap", "3d")))
            if self._uruchom("rap", [py, "-u", str(TU / "40_widok_3d.py"),
                                     str(self.wynik_folder)]) != 0:
                if self.przerwane:
                    return False, ZATRZYMANO
                self.q.put(("info", T("Widok 3D się nie zrobił - reszta "
                                      "wyników jest w porządku.\n",
                                      "The 3D view was not created - the "
                                      "other results are fine.\n")))
        self.q.put(("faza_koniec", "rap"))
        return True, T("Gotowe.", "Done.")

    # ------------------------------------------------------------------
    # Kolejka na noc: wiele skanów po kolei, jeden zestaw w Excelu.
    def pokaz_kolejke(self, dodaj=None):
        if self.okno_kolejki and self.okno_kolejki.winfo_exists():
            self.okno_kolejki.lift()
        else:
            self.okno_kolejki = OknoKolejki(self)
        if dodaj:
            self.okno_kolejki.dodaj(dodaj)

    def start_kolejki(self, po_koncu, pomijaj):
        if not self.kolejka:
            return
        self.w_kolejce = True
        self.przerwane = False
        self.po_koncu = po_koncu
        self.pomijaj = pomijaj
        for poz in self.kolejka:
            if poz["stan"] != "gotowe" or not pomijaj:
                poz.update(stan="czeka", blad="", czas=None, v=None)
        self.plik_zest = WYNIKI / (T("KOLEJKA_", "QUEUE_") +
                                   f"{time.strftime('%Y-%m-%d_%H%M')}.xlsx")
        self.b_start.config(state="disabled")
        self.b_kolejka.config(state="normal")
        self.b_stop.config(state="normal")
        self.t_kolejki = time.time()
        threading.Thread(target=self._praca_kolejka, daemon=True).start()

    def _praca_kolejka(self):
        py = interpreter()
        czuwaj(True)
        try:
            if not self._biblioteki(py):
                return self.q.put(("kolejka_koniec", False, BLAD_LIB))
            for i, poz in enumerate(self.kolejka):
                if self.przerwane:
                    break
                if poz["stan"] == "gotowe":
                    continue
                self._ustaw_skan(poz["skan"])
                poz["stan"] = "liczy się"
                self.ost_v = None
                self.ost_blad = ""
                self.q.put(("nowy_skan", i))
                t0 = time.time()
                try:
                    if not self.skan.is_file():
                        ok, tekst = False, T("Nie ma pliku skanu.",
                                             "The scan file is missing.")
                    else:
                        zmiana = self._zapisz_ustawienia(poz["klasa"],
                                                         poz["gatunek"])
                        ok, tekst = self._licz_skan(
                            py, pomijaj=self.pomijaj and not zmiana)
                except Exception as e:               # noqa: BLE001
                    ok, tekst = False, T(f"Błąd: {e}", f"Error: {e}")
                poz.update(stan="gotowe" if ok else
                           ("zatrzymano" if self.przerwane else "błąd"),
                           blad="" if ok else " ".join(
                               x for x in (tekst, self.ost_blad) if x),
                           czas=time.time() - t0,
                           folder=self.wynik_folder)
                self.q.put(("skan_koniec", i, ok, tekst, self.wynik_folder))
                self._zestawienie()
            self._zestawienie()
            n_ok = sum(p["stan"] == "gotowe" for p in self.kolejka)
            self.q.put(("kolejka_koniec", not self.przerwane, T(
                f"Kolejka skończona: {n_ok} z {len(self.kolejka)} skanów "
                f"policzonych.", f"Queue finished: {n_ok} of "
                f"{len(self.kolejka)} scans processed.")))
        except Exception as e:                       # noqa: BLE001
            self.q.put(("kolejka_koniec", False, T(f"Błąd kolejki: {e}",
                                                   f"Queue error: {e}")))
        finally:
            czuwaj(False)

    def _zestawienie(self):
        """Excel z wszystkimi powierzchniami z kolejki (zapis po każdym
        skanie, więc nawet po awarii coś zostaje)."""
        kol = [T("Lp.", "No."), T("Skan", "Scan"), T("Stan", "Status"),
               T("Czas [min]", "Time [min]"), T("Drzewostan", "Stand"),
               T("Gatunek", "Species"),
               T("Powierzchnia [m²]", "Plot area [m²]"),
               T("Liczba drzew", "Number of trees"),
               T("N [szt./ha]", "N [trees/ha]"),
               "Dg [cm]", T("Wysokość Loreya [m]", "Lorey's height [m]"),
               "G [m²/ha]",
               T("Miąższość – powierzchnia [m³]", "Volume – plot [m³]"),
               "V [m³/ha]", T("Krzewy pominięte", "Shrubs excluded"),
               T("Złomy", "Snags"), T("Uwagi / błąd", "Notes / error"),
               T("Folder wyników", "Results folder")]
        wiersze = []
        for i, poz in enumerate(self.kolejka, 1):
            folder = poz.get("folder") or WYNIKI / Path(poz["skan"]).stem
            ps = podsumowanie(folder) if poz["stan"] == "gotowe" else {}
            war = []
            for k_pl, k_en in KLUCZE_PODS:
                v_ = wartosc_pods(ps, k_pl)
                war.append(v_ if v_ is not None else wartosc_pods(ps, k_en))
            if war[7] is None and poz["stan"] == "gotowe":
                war[7] = self.ost_v
            wiersze.append([i, Path(poz["skan"]).stem,
                            STAN.get(poz["stan"], poz["stan"]),
                            round(poz["czas"] / 60, 1) if poz.get("czas")
                            else None, KROTKO_KLASA.get(poz["klasa"]),
                            KROTKO_GAT.get(poz["gatunek"])] +
                           war[:8] + war[8:] + [poz.get("blad", ""),
                                                str(folder)])
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = T("Kolejka", "Queue")
            ws.append([T(f"{J.NAZWA} – zestawienie kolejki, ",
                         f"{J.NAZWA} – queue summary, ") +
                       f"{time.strftime('%Y-%m-%d %H:%M')}"])
            ws["A1"].font = Font(bold=True, size=13)
            ws.append([])
            ws.append(kol)
            for c in ws[3]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="4E7A3A")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for w_ in wiersze:
                ws.append(w_)
                if w_[2] != STAN["gotowe"]:
                    for c in ws[ws.max_row]:
                        c.font = Font(color="B03030")
            ok = [w_ for w_ in wiersze if w_[2] == STAN["gotowe"]]
            if ok:
                ws.append([])
                sr = ["", T("Średnio (policzone)", "Mean (processed)"), "", "",
                      "", ""]
                for j in range(6, 16):
                    v = [w_[j] for w_ in ok if w_[j] is not None]
                    sr.append(round(sum(v) / len(v), 1) if v else None)
                ws.append(sr)
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
            for j, szer in enumerate([5, 22, 11, 9, 13, 14, 12, 9, 9, 8, 10,
                                      9, 13, 9, 9, 7, 40, 50], 1):
                ws.column_dimensions[ws.cell(3, j).column_letter].width = szer
            ws.row_dimensions[3].height = 32
            ws.freeze_panes = "C4"
            WYNIKI.mkdir(parents=True, exist_ok=True)
            wb.save(self.plik_zest)
        except Exception:                            # noqa: BLE001
            try:                                 # bez openpyxl: zwykły CSV
                with open(self.plik_zest.with_suffix(".csv"), "w",
                          encoding="utf-8-sig") as f:
                    f.write(";".join(kol) + "\n")
                    for w_ in wiersze:
                        f.write(";".join("" if x is None else
                                         T(str(x).replace(".", ","), str(x))
                                         if isinstance(x, float) else str(x)
                                         for x in w_) + "\n")
            except OSError:
                pass

    # ------------------------------------------------------------------
    # Petla okna: odbiera wiadomosci z watku pracy.
    def _petla(self):
        try:
            while True:
                self._obsluz(self.q.get_nowait())
        except queue.Empty:
            pass
        if self.t_start and str(self.b_stop["state"]) == "normal":
            s = int(time.time() - self.t_start)
            self.l_czas.config(text=f"{s // 60:02d}:{s % 60:02d}")
            self._pulsuj()
        self._odswiez_obrazek()
        self._odswiez_zywo()
        self.root.after(150, self._petla)

    def _indeks(self, faza, krok):
        for i, (f, k, _, _) in enumerate(ETAPY):
            if f == faza and k == krok:
                return i
        if faza == "drz" and krok == 7:              # KROK 7 i 8 to jeden etap
            return self._indeks("drz", 8)
        return None

    def _ustaw(self, i):
        """Etap i staje sie biezacy, wszystkie wczesniejsze sa zrobione."""
        for j in range(i):
            if self.stan[j] != "pominiety":
                self.stan[j] = "zrobiony"
                self.ulamek[j] = 1.0
        self.stan[i] = "trwa"
        self.biezacy = i
        self._rysuj_etapy()

    def _obsluz(self, msg):
        typ = msg[0]
        if typ == "etap":
            self._ustaw(msg[1])
        elif typ == "gotowe":
            self.stan[msg[1]] = "zrobiony"; self.ulamek[msg[1]] = 1.0
            self._rysuj_etapy()
        elif typ == "pomin":
            for i, (f, *_r) in enumerate(ETAPY):
                if f == msg[1]:
                    self.stan[i] = "pominiety"; self.ulamek[i] = 1.0
            self._rysuj_etapy()
            if msg[1] == "drz":
                self._policz_drzewa()
        elif typ == "faza_koniec":
            for i, (f, *_r) in enumerate(ETAPY):
                if f == msg[1] and self.stan[i] != "pominiety":
                    self.stan[i] = "zrobiony"; self.ulamek[i] = 1.0
            self._rysuj_etapy()
            if msg[1] == "drz":
                self._policz_drzewa()
        elif typ == "info":
            self._dopisz(msg[1], "info")
        elif typ == "status":
            self.l_status.config(text=msg[1])
        elif typ == "linia":
            self._linia(msg[1], msg[2])
        elif typ == "koniec":
            self._koniec(msg[1], msg[2])
        elif typ == "nowy_skan":
            i = msg[1]
            self._wyczysc()
            self.t_start = time.time()
            self.var_skan.set(str(self.kolejka[i]["skan"]))
            self.l_kolejka.config(text=T(f"Kolejka: skan {i + 1} z ",
                                         f"Queue: scan {i + 1} of ") +
                                  f"{len(self.kolejka)}")
            self.root.title(f"{J.NAZWA} – " + T("kolejka", "queue") +
                            f" {i + 1}/{len(self.kolejka)}: "
                            f"{Path(self.kolejka[i]['skan']).stem}")
            self._dopisz(f"===== {Path(self.kolejka[i]['skan']).name} "
                         f"=====\n", "info")
            self._odswiez_kolejke()
        elif typ == "skan_koniec":
            _, i, ok, tekst, folder = msg
            if not ok and self.biezacy is not None:
                self.stan[self.biezacy] = "blad"
                self._rysuj_etapy()
            self._dopisz(f"\n{T('GOTOWE', 'DONE') if ok else tekst}: "
                         f"{Path(self.kolejka[i]['skan']).name}\n",
                         "info" if ok else "blad")
            self._zapisz_dziennik(folder)
            self._odswiez_kolejke()
        elif typ == "kolejka_koniec":
            self._koniec_kolejki(msg[1], msg[2])

    def _linia(self, faza, linia):
        blad = ("Traceback" in linia or "Error" in linia or
                linia.startswith(("BLAD", "ERROR")))
        self._dopisz(linia, "blad" if blad else None)
        m = WZOR_KROK.search(linia)
        if m and faza in ("drz", "cz"):
            i = self._indeks(faza, int(m.group(1)))
            if i is not None and (self.biezacy is None or i >= self.biezacy):
                self._ustaw(i)
                self.l_status.config(text=ETAPY[i][2] + "…")
        if faza == "cz" and linia.startswith("[GPKG]"):
            i = self._indeks("cz", "gpkg")
            self._ustaw(i)
            self.l_status.config(text=ETAPY[i][2] + "…")
        m = WZOR_RYS.search(linia)
        if m:
            i = self._indeks("cz", "img")
            n, razem = int(m.group(1)), int(m.group(2))
            if self.biezacy != i:
                self._ustaw(i)
            self.ulamek[i] = n / razem
            self.wiersze[i][2].config(text=f"{n} / {razem}")
            self.l_status.config(text=T(f"Rysuję drzewo {n} z {razem}…",
                                        f"Drawing tree {n} of {razem}…"))
        m = WZOR_PROC.match(linia)
        if m:
            klucz = {"pien": "pień", "galezie": "gałęzie",
                     "ulistnienie": "ulistnienie", "stem": "pień",
                     "branches": "gałęzie",
                     "foliage": "ulistnienie"}[m.group(1)]
            self.kafle[klucz].config(text=f"{float(m.group(2)):.0f}%")
        m = WZOR_V.search(linia)
        if m:
            self.kafle["m³/ha"].config(text=f"{float(m.group(1)):.0f}")
        m = WZOR_DRZEW.search(linia)
        if m:
            self.kafle["drzew"].config(text=m.group(1))
            i = self._indeks("drz", 5)
            self.wiersze[i][2].config(text=T(f"{m.group(1)} pni",
                                             f"{m.group(1)} stems"))

    def _policz_drzewa(self):
        csv = U.plik(self.wynik_folder, f"{self.nazwa}_drzewa.csv")
        try:
            n = sum(1 for _ in open(csv, encoding="utf-8")) - 1
            self.kafle["drzew"].config(text=str(n))
        except OSError:
            pass

    def _dopisz(self, tekst, tag=None):
        self.dziennik_caly.append(tekst)
        self.log.config(state="normal")
        self.log.insert("end", tekst, tag or ())
        if int(self.log.index("end-1c").split(".")[0]) > 3000:
            self.log.delete("1.0", "500.0")
        self.log.see("end")
        self.log.config(state="disabled")

    def _rysuj_etapy(self):
        for (ikona, napis, dod), st in zip(self.wiersze, self.stan):
            if st == "zrobiony":
                ikona.config(text="✓", fg=AKCENT); napis.config(fg=TEKST)
            elif st == "trwa":
                ikona.config(text="●", fg=DREWNO); napis.config(fg=TEKST)
            elif st == "pominiety":
                ikona.config(text="↷", fg=SZARY); napis.config(fg=SZARY)
                dod.config(text=T("było wcześniej", "done earlier"))
            elif st == "blad":
                ikona.config(text="✕", fg=BLAD); napis.config(fg=BLAD)
            else:
                ikona.config(text="○", fg=SZARY); napis.config(fg=SZARY)
        razem = sum(e[3] for e in ETAPY)
        zrob = sum(e[3] * u for e, u in zip(ETAPY, self.ulamek))
        self.pasek["value"] = 1000 * zrob / razem
        self.l_proc.config(text=f"{100 * zrob / razem:.0f}%")

    def _pulsuj(self):
        if self.biezacy is not None and self.stan[self.biezacy] == "trwa":
            ikona = self.wiersze[self.biezacy][0]
            faza = int(time.time() * 2) % 2
            ikona.config(fg=DREWNO if faza else "#7a5528")
            self._rysuj_pasek_tylko()

    def _rysuj_pasek_tylko(self):
        razem = sum(e[3] for e in ETAPY)
        zrob = sum(e[3] * u for e, u in zip(ETAPY, self.ulamek))
        self.pasek["value"] = 1000 * zrob / razem
        self.l_proc.config(text=f"{100 * zrob / razem:.0f}%")

    def _odswiez_obrazek(self):
        if not self.wynik_folder:
            return
        if getattr(self, "_t_obr", 0) > time.time() - 1.0:
            return
        self._t_obr = time.time()
        kat = self.wynik_folder / U.OBRAZY
        if not kat.exists():
            return
        try:
            pliki = sorted(kat.glob("drzewo_*.png"))
        except OSError:
            return
        if not pliki or pliki[-1] == self.ostatni_obrazek:
            return
        # tylko pliki z tego przebiegu
        if self.t_start and pliki[-1].stat().st_mtime < self.t_start:
            return
        p = pliki[-1]
        try:
            img = tk.PhotoImage(file=str(p))
        except tk.TclError:
            return                                   # plik jeszcze sie zapisuje
        h = max(1, self.platno.winfo_height() - 10)
        skala = max(1, -(-img.height() // h))
        if skala > 1:
            img = img.subsample(skala, skala)
        if self.ostatni_obrazek is None:      # pierwszy obrazek tego przebiegu
            self.zakladki.select(1)
        self.obrazek = img
        self.ostatni_obrazek = p
        self.platno.config(image=img, text="")
        self.l_obrazek.config(text=p.name)
        self.b_podglad.config(state="normal")

    # ------------------------------------------------------------------
    # Podgląd na żywo: programy liczące zapisują na_zywo.npz, tu rysujemy.
    def _odswiez_zywo(self):
        p = self.plik_zywo
        if not p or time.time() - getattr(self, "_t_spr_zywo", 0) < 1.0:
            return
        self._t_spr_zywo = time.time()
        try:
            t = p.stat().st_mtime
        except OSError:
            return
        if t <= self.t_zywo or (self.t_start and t < self.t_start):
            return
        try:
            import numpy as np
            with np.load(p) as d:
                dane = {k: d[k] for k in d.files}
        except Exception:                                # noqa: BLE001
            return                                       # jeszcze się zapisuje
        self.t_zywo = t
        try:
            self._rysuj_zywo(dane)
        except Exception as e:                           # noqa: BLE001
            self.l_zywo_pusty.config(text=T(f"Podgląd niedostępny: {e}",
                                            f"Preview not available: {e}"))
            return
        if self.ostatni_obrazek is None:
            self.zakladki.select(0)

    def _rysuj_zywo(self, d):
        import numpy as np
        if self.zywo is None:
            import matplotlib
            matplotlib.use("TkAgg")
            from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
            from matplotlib.figure import Figure
            fig = Figure(facecolor=PANEL)
            gs = fig.add_gridspec(1, 2, width_ratios=[1.15, 1], wspace=0.08,
                                  left=0.02, right=0.98, top=0.84, bottom=0.03)
            osie = (fig.add_subplot(gs[0]), fig.add_subplot(gs[1]))
            plotno = FigureCanvasTkAgg(fig, master=self.f_zywo)
            self.l_zywo_pusty.pack_forget()
            plotno.get_tk_widget().pack(fill="both", expand=True)
            self.zywo = (fig, plotno, osie)
        fig, plotno, (a1, a2) = self.zywo
        xyz, w, rodzaj = d["xyz"], d["w"], str(d["rodzaj"])
        if rodzaj == "czesci":
            kol = np.array(KOLOR_CZESCI)[np.clip(w.astype(int), 0, 4)]
        elif rodzaj == "drzewa":
            from matplotlib import colormaps
            paleta = colormaps["tab20"](np.arange(20))
            kol = paleta[(w.astype(int) * 7) % 20]
            kol[w <= 0] = (0.29, 0.33, 0.31, 1.0)
        else:
            from matplotlib import colormaps
            gora = max(float(np.percentile(w, 99)), 1.0) if len(w) else 1.0
            kol = colormaps["viridis"](np.clip(w / gora, 0, 1))
        for a in (a1, a2):
            a.clear()
            a.set_facecolor("#0f1214")
            a.set_xticks([]); a.set_yticks([])
            for s_ in a.spines.values():
                s_.set_color(LINIA)
        if len(xyz):
            # z góry: najpierw niskie punkty, wysokie na wierzchu
            k = np.argsort(xyz[:, 2])
            a1.scatter(xyz[k, 0], xyz[k, 1], c=kol[k], s=0.4, linewidths=0)
            a1.set_aspect("equal", adjustable="datalim")
            # z boku: pas przez środek chmury, żeby było widać pnie
            sy = np.median(xyz[:, 1])
            pas = np.abs(xyz[:, 1] - sy) < max(2.0, np.ptp(xyz[:, 1]) * 0.15)
            a2.scatter(xyz[pas, 0], xyz[pas, 2], c=kol[pas], s=0.4,
                       linewidths=0)
        for x, y, r in d["okregi"]:
            a1.add_patch(_okrag(x, y, max(r, 0.12)))
        a1.set_title(T("z góry", "top view"), color=SZARY, fontsize=9)
        a2.set_title(T("z boku (pas przez środek)",
                       "side view (strip through the centre)"), color=SZARY,
                     fontsize=9)
        ile = T(f"{int(d['punktow']):,}".replace(",", " "),
                f"{int(d['punktow']):,}")
        fig.suptitle(f"{d['opis']}   ·   {ile} " + T("punktów", "points"),
                     color=TEKST, fontsize=10, y=0.985)
        plotno.draw_idle()

    def _odswiez_kolejke(self):
        for poz in self.kolejka:
            if poz["stan"] == "gotowe" and poz.get("v") is None:
                ps = podsumowanie(poz.get("folder") or
                                  WYNIKI / Path(poz["skan"]).stem)
                poz["v"] = ps.get("Miąższość grubizny (V)",
                                  ps.get("Merchantable volume ≥7 cm (V)"))
        if self.okno_kolejki and self.okno_kolejki.winfo_exists():
            self.okno_kolejki.odswiez()

    def _koniec_kolejki(self, ok, tekst):
        self.w_kolejce = False
        self.b_start.config(state="normal")
        self.b_kolejka.config(state="normal")
        self.b_stop.config(state="disabled")
        self.root.title(J.NAZWA)
        self.l_kolejka.config(text="")
        self._odswiez_kolejke()
        s = int(time.time() - self.t_kolejki)
        zest = self.plik_zest if self.plik_zest.exists() else \
            self.plik_zest.with_suffix(".csv")
        self.l_status.config(text=T(f"{tekst} Czas {s // 3600} h "
                                    f"{s % 3600 // 60} min. Zestawienie: "
                                    f"{zest.name}", f"{tekst} Time "
                                    f"{s // 3600} h {s % 3600 // 60} min. "
                                    f"Summary: {zest.name}"))
        self._dopisz(T(f"\n{tekst}\nZestawienie: {zest}\n",
                       f"\n{tekst}\nSummary: {zest}\n"),
                     "info" if ok else "blad")
        gotowe = [p for p in self.kolejka if p["stan"] == "gotowe"]
        if gotowe:                    # przyciski do ostatniego udanego skanu
            self._ustaw_skan(gotowe[-1]["skan"])
            self.var_skan.set(str(self.skan))
        if self.wynik_folder:
            self._przyciski(self.wynik_folder, self.nazwa)
        if ok and self.po_koncu and os.name == "nt":
            self._odliczaj(self.po_koncu, 60)

    def _odliczaj(self, co, sek):
        """Ostatnia minuta przed uśpieniem / wyłączeniem - da się anulować."""
        okno = tk.Toplevel(self.root)
        okno.title(J.NAZWA)
        okno.configure(bg=TLO)
        okno.attributes("-topmost", True)
        napis = ttk.Label(okno, style="Pod.TLabel", padding=24)
        napis.pack()
        stan = {"s": sek, "anuluj": False}
        ttk.Button(okno, text=T("Anuluj", "Cancel"), command=lambda: (
            stan.update(anuluj=True), okno.destroy())).pack(pady=(0, 18))

        def tik():
            if stan["anuluj"]:
                return
            if stan["s"] <= 0:
                okno.destroy()
                flagi = subprocess.CREATE_NO_WINDOW
                if co == "wyl":
                    subprocess.run(["shutdown", "/s", "/t", "0"],
                                   creationflags=flagi)
                else:
                    subprocess.run(["rundll32.exe", "powrprof.dll,"
                                    "SetSuspendState", "0,1,0"],
                                   creationflags=flagi)
                return
            napis.config(text=T(
                f"Kolejka skończona. Komputer "
                f"{'wyłączy się' if co == 'wyl' else 'uśpi się'}"
                f" za {stan['s']} s.", f"Queue finished. The computer will "
                f"{'shut down' if co == 'wyl' else 'go to sleep'} in "
                f"{stan['s']} s."))
            stan["s"] -= 1
            okno.after(1000, tik)
        tik()

    def _zapisz_dziennik(self, folder=None):
        try:
            kat = U.katalog(folder or self.wynik_folder, U.DZIENNIK)
            (kat / f"dziennik_{time.strftime('%Y-%m-%d_%H%M')}.txt").write_text(
                "".join(self.dziennik_caly), encoding="utf-8")
        except OSError:
            pass

    def _koniec(self, ok, tekst):
        self.b_start.config(state="normal")
        self.b_kolejka.config(state="normal")
        self.b_stop.config(state="disabled")
        if self.wynik_folder and self.wynik_folder.exists():
            self._przyciski(self.wynik_folder, self.nazwa)
        if ok:
            self.stan = ["pominiety" if s == "pominiety" else "zrobiony"
                         for s in self.stan]
            self.ulamek = [1.0] * len(ETAPY)
            self._rysuj_etapy()
            s = int(time.time() - self.t_start)
            self.l_status.config(text=T(
                f"Gotowe w {s // 60} min {s % 60} s. Wszystko jest w "
                f"{self.wynik_folder}", f"Done in {s // 60} min {s % 60} s. "
                f"Everything is in {self.wynik_folder}"))
            self._dopisz(T(f"\nGOTOWE. Wyniki: {self.wynik_folder}\n",
                           f"\nDONE. Results: {self.wynik_folder}\n"),
                         "info")
        else:
            if self.biezacy is not None:
                self.stan[self.biezacy] = "blad"
                self._rysuj_etapy()
            self.l_status.config(text=T(
                f"{tekst}  Kliknij „Kopiuj dziennik” i wklej go Claude'owi.",
                f"{tekst}  Click “Copy log” and paste it to Claude."))
            self._dopisz(f"\n{tekst}\n", "blad")
        if self.wynik_folder and self.wynik_folder.exists():
            self._zapisz_dziennik()


class OknoKolejki(tk.Toplevel):
    """Lista skanów do policzenia po kolei (np. przez noc)."""

    def __init__(self, okno):
        super().__init__(okno.root)
        self.o = okno
        self.title(J.NAZWA + T(" – kolejka na noc", " – overnight queue"))
        self.configure(bg=TLO)
        self.geometry("1010x560")
        self.minsize(760, 420)
        s = ttk.Style()
        s.configure("Kolejka.Treeview", background=PANEL, fieldbackground=PANEL,
                    foreground=TEKST, rowheight=24, bordercolor=LINIA)
        s.configure("Kolejka.Treeview.Heading", background=LINIA,
                    foreground=TEKST, relief="flat")
        s.map("Kolejka.Treeview", background=[("selected", "#3d5a30")])

        ttk.Label(self, text=T("Dodaj skany, ustaw drzewostan i gatunek, "
                               "kliknij Start. Skany liczą się po kolei; błąd "
                               "jednego nie zatrzymuje reszty.",
                               "Add scans, set the stand and species, click "
                               "Start. The scans are processed one after "
                               "another; an error in one does not stop the "
                               "rest."), style="Pod.TLabel",
                  padding=(16, 12, 16, 6), wraplength=860).pack(fill="x")
        g = ttk.Frame(self, padding=(16, 0))
        g.pack(fill="x")
        ttk.Button(g, text=T("Dodaj pliki…", "Add files…"),
                   command=self.dodaj_pliki).pack(side="left")
        ttk.Button(g, text=T("Dodaj folder…", "Add folder…"),
                   command=self.dodaj_folder).pack(side="left", padx=(8, 0))
        ttk.Button(g, text=T("Usuń zaznaczone", "Remove selected"),
                   command=self.usun).pack(
            side="left", padx=(8, 0))
        ttk.Button(g, text="↑", width=3, command=lambda: self.przesun(-1)
                   ).pack(side="left", padx=(8, 0))
        ttk.Button(g, text="↓", width=3, command=lambda: self.przesun(1)
                   ).pack(side="left", padx=(4, 0))

        u = ttk.Frame(self, padding=(16, 8, 16, 0))
        u.pack(fill="x")
        ttk.Label(u, text=T("Zaznaczonym ustaw:  drzewostan",
                            "For selected set:  stand"),
                  style="Pod.TLabel").pack(side="left")
        self.var_k = tk.StringVar(value=okno.var_klasa.get())
        ttk.Combobox(u, textvariable=self.var_k, values=list(okno.KLASY),
                     state="readonly", width=28).pack(side="left", padx=6)
        ttk.Label(u, text=T("gatunek", "species"), style="Pod.TLabel").pack(
            side="left")
        self.var_g = tk.StringVar(value=okno.var_gatunek.get())
        ttk.Combobox(u, textvariable=self.var_g, values=list(okno.GATUNKI),
                     state="readonly", width=26).pack(side="left", padx=6)
        ttk.Button(u, text=T("Ustaw", "Apply"), command=self.ustaw).pack(
            side="left")

        dol = ttk.Frame(self, padding=(16, 8, 16, 14))
        dol.pack(side="bottom", fill="x")
        self.var_pomijaj = tk.BooleanVar(value=True)
        tk.Checkbutton(dol, text=T("Pomijaj etapy już policzone",
                                   "Skip stages already computed"),
                       variable=self.var_pomijaj, bg=TLO, fg=TEKST,
                       selectcolor=PANEL, activebackground=TLO,
                       activeforeground=TEKST, font=("Segoe UI", 10)
                       ).pack(side="left")
        ttk.Label(dol, text=T("   Po zakończeniu:", "   When finished:"),
                  style="Pod.TLabel").pack(
            side="left")
        self.var_po = tk.StringVar(value=next(iter(PO_KOLEJCE)))
        ttk.Combobox(dol, textvariable=self.var_po, values=list(PO_KOLEJCE),
                     state="readonly", width=16).pack(side="left", padx=6)
        self.b_start = ttk.Button(dol, text=T("▶  Start kolejki",
                                              "▶  Start queue"),
                                  style="Akcent.TButton", command=self.start)
        self.b_start.pack(side="right")
        self.l_info = ttk.Label(dol, text="", style="Pod.TLabel")
        self.l_info.pack(side="right", padx=10)

        t = ttk.Frame(self, padding=(16, 8))
        t.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(t, style="Kolejka.Treeview",
                                 columns=("skan", "klasa", "gatunek", "stan",
                                          "czas", "v"), show="headings",
                                 selectmode="extended")
        for c, n, w in (("skan", T("Skan", "Scan"), 240),
                        ("klasa", T("Drzewostan", "Stand"), 110),
                        ("gatunek", T("Gatunek", "Species"), 110),
                        ("stan", T("Stan", "Status"), 110),
                        ("czas", T("Czas", "Time"), 70), ("v", "m³/ha", 70)):
            self.tree.heading(c, text=n)
            self.tree.column(c, width=w, anchor="w" if c == "skan"
                             else "center")
        self.tree.tag_configure("gotowe", foreground=AKCENT)
        self.tree.tag_configure("błąd", foreground=BLAD)
        self.tree.tag_configure("liczy się", foreground=DREWNO)
        sb = ttk.Scrollbar(t, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="left", fill="y")
        self.odswiez()

    def dodaj(self, sciezki):
        znane = {str(Path(p["skan"]).resolve()).lower() for p in self.o.kolejka}
        for f in skany_w(sciezki):
            if str(f.resolve()).lower() in znane:
                continue
            u = UST.wczytaj(WYNIKI / f.stem, f.stem)
            self.o.kolejka.append({
                "skan": f,
                "klasa": u.get("klasa", self.o.KLASY[self.var_k.get()]),
                "gatunek": u.get("gatunek", self.o.GATUNKI[self.var_g.get()]),
                "stan": "czeka", "blad": "", "czas": None})
        self.odswiez()

    def dodaj_pliki(self):
        p = filedialog.askopenfilenames(
            parent=self, title=T("Wybierz skany", "Choose scans"),
            filetypes=[(T("Chmury punktów", "Point clouds"), "*.laz *.las")],
            initialdir=str(ROOT.parent))
        if p:
            self.dodaj(p)

    def dodaj_folder(self):
        p = filedialog.askdirectory(parent=self, title=T("Folder ze skanami",
                                                         "Folder with scans"),
                                    initialdir=str(ROOT.parent))
        if p:
            self.dodaj([p])

    def _zaznaczone(self):
        return sorted(int(i) for i in self.tree.selection())

    def usun(self):
        if self.o.w_kolejce:
            return
        for i in reversed(self._zaznaczone()):
            del self.o.kolejka[i]
        self.odswiez()

    def przesun(self, k):
        if self.o.w_kolejce:
            return
        zaz = self._zaznaczone()
        kol = self.o.kolejka
        for i in (zaz if k < 0 else reversed(zaz)):
            j = i + k
            if 0 <= j < len(kol):
                kol[i], kol[j] = kol[j], kol[i]
        self.odswiez([i + k for i in zaz if 0 <= i + k < len(kol)])

    def ustaw(self):
        for i in self._zaznaczone():
            if self.o.kolejka[i]["stan"] != "liczy się":
                self.o.kolejka[i]["klasa"] = self.o.KLASY[self.var_k.get()]
                self.o.kolejka[i]["gatunek"] = \
                    self.o.GATUNKI[self.var_g.get()]
        self.odswiez(self._zaznaczone())

    def odswiez(self, zaznacz=None):
        self.tree.delete(*self.tree.get_children())
        for i, p in enumerate(self.o.kolejka):
            cz = p.get("czas")
            v = p.get("v")
            self.tree.insert("", "end", iid=str(i), tags=(p["stan"],), values=(
                Path(p["skan"]).name, KROTKO_KLASA.get(p["klasa"], p["klasa"]),
                KROTKO_GAT.get(p["gatunek"], p["gatunek"]),
                STAN.get(p["stan"], p["stan"]),
                (f"{int(cz // 60)} min" if cz >= 60 else "<1 min")
                if cz is not None else "",
                f"{v:.0f}" if v is not None else ""))
        if zaznacz:
            self.tree.selection_set([str(i) for i in zaznacz])
        n = len(self.o.kolejka)
        gotowe = sum(p["stan"] == "gotowe" for p in self.o.kolejka)
        self.l_info.config(text=T(f"{n} skanów", f"{n} scan" + ("s" if n != 1 else "")) + (
            T(f", gotowe {gotowe}", f", done {gotowe}") if gotowe else ""))
        self.b_start.config(state="disabled" if self.o.w_kolejce or not n
                            or str(self.o.b_stop["state"]) == "normal"
                            else "normal")

    def start(self):
        if not self.o.kolejka:
            return
        self.o.start_kolejki(PO_KOLEJCE[self.var_po.get()],
                             self.var_pomijaj.get())
        self.odswiez()


def _okrag(x, y, r):
    from matplotlib.patches import Circle
    return Circle((x, y), r, fill=False, ec="#ff6b5a", lw=1.0)


def main():
    root = tk.Tk()
    try:
        root.tk.call("tk", "scaling", root.winfo_fpixels("1i") / 72.0)
    except tk.TclError:
        pass
    Okno(root, sys.argv[1:] or None)
    root.mainloop()


if __name__ == "__main__":
    main()
