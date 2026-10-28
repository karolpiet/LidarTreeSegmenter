"""
LidarTreeSegmenter - krok 3: RAPORT (Excel + podsumowanie tekstowe)
======================================================

Zbiera wszystko, co powstalo dla jednego skanu, w jeden porzadny plik
Excela i w czytelne 00_PODSUMOWANIE.txt:

  <nazwa>_RAPORT.xlsx   arkusze: Podsumowanie, Drzewa, Ustawienia,
                        Objaśnienia
  00_PODSUMOWANIE.txt   podsumowanie drzewostanu (z miąższością)
                        i ustawienia opisane pełnymi zdaniami

Uzycie (zwykle robi to samo okno LidarTreeSegmenter albo PRZECIAGNIJ_TU_SKAN.bat):

  python 30_raport.py wyniki\\<nazwa_skanu>
  python 30_raport.py wyniki\\<nazwa_skanu> --grupa iglaste

MIĄŻSZOŚĆ
---------
Dwie wartosci, bo kazda mowi co innego:

1. Miąższość grubizny - wzór zależy od gatunku drzewa i jest brany
   z pliku tablice_miazszosci.csv (obok programu). Domyślnie są tam
   wzory ogólne V = g * h * f (g - pole przekroju w pierśnicy,
   h - wysokość, f - pierśnicowa liczba kształtu grubizny): świerk
   0,47, buk 0,51, nieokreślony 0,47. Plik można podmienić na
   tablice regionalne, także wzory potęgowe V = a * d^b * h^c.
   Gatunek drzewa bierzemy z <skan>_gatunki.csv (kolumna
   gatunek_reczny, wypełniasz sam), a gdy pusta - rozpoznany
   automatycznie (iglaste = świerk, liściaste = buk), a gdy brak -
   domyślny (--grupa).

2. Miąższość zmierzonego odcinka strzały - suma walców ściętych między
   kolejnymi okręgami dopasowanymi do pnia przez LidarTreeSegmenter (krok 10).
   Obejmuje TYLKO odcinek, który skaner widział (plus dorysowany dół).
   Nie zawiera wierzchołka, więc zawsze jest mniejsza od grubizny.
   Przydaje się do kontroli: jeśli jest WIĘKSZA od grubizny, coś jest
   nie tak z pomiarem tego drzewa.
"""

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

import lts_jezyk as J
from lts_jezyk import t

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print(t("BRAK BIBLIOTEKI openpyxl - zainstaluj: ",
            "MISSING LIBRARY openpyxl - install it: ") +
          "python -m pip install openpyxl")
    sys.exit(1)

TU = Path(__file__).resolve().parent       # LidarTreeSegmenter\\program
ROOT = TU.parent
TABLICE = ROOT / "tablice_miazszosci.csv"
sys.path.insert(0, str(TU))
import lts_uklad as U                    # noqa: E402
import lts_ustawienia as UST             # noqa: E402
import lts_sciezka as SC                 # noqa: E402
import lts_dtm as DTM                    # noqa: E402
import lts_etykiety as ET                # noqa: E402
import lts_krzewy as KR                  # noqa: E402
import lts_walec as LW                   # noqa: E402
GRUPA_DOMYSLNA = "nieokreślony"
# gatunek bez własnego wiersza w tablicach liczy się wzorem swojej grupy
IGLASTE = {"świerk", "jodła", "sosna", "modrzew", "daglezja", "cis"}


def wiersz_tablic(gr, tab):
    """Wiersz tablic dla gatunku: własny, grupy (iglaste / liściaste) albo
    domyślny."""
    if gr in tab.index:
        return gr
    grupa = "iglaste" if gr in IGLASTE else (
        "liściaste" if gr in J.GATUNKI_EN and gr not in (
            "nieokreślony", "nieokreślone", "inny", "inne", "iglaste",
            "liściaste") else None)
    if grupa in tab.index:
        return grupa
    return GRUPA_DOMYSLNA if GRUPA_DOMYSLNA in tab.index else tab.index[0]
# główne gatunki na mierzonych powierzchniach: świerk i buk
GATUNEK_GRUPY = {"iglaste": "świerk", "liściaste": "buk",
                 "nieokreślone": "nieokreślony"}

# kolory raportu
GRANAT = "1F3B57"
JASNY = "EAF1F8"
NIEBIESKI = "D6E6FA"     # wartosci dorysowane (rekonstrukcja)
ZOLTY = "FFF2CC"         # wartosci szacunkowe / z uwagami
SZARY = "808080"


# ----------------------------------------------------------------------
# liczby po polsku
# ----------------------------------------------------------------------
def pl(x, miejsc=1):
    """Liczba po polsku (1 234,5); w wersji angielskiej 1,234.5."""
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "–"
    return J.liczba(x, miejsc)


# ----------------------------------------------------------------------
# ustawienia z kroku 1 (tree_segmenter_tool)
# ----------------------------------------------------------------------
def wczytaj_ustawienia(katalog, nazwa):
    """Nowa wersja kroku 1 zapisuje <nazwa>_ustawienia.json. Starsze
    wyniki maja tylko 00_PODSUMOWANIE.txt - wtedy czytam wartosci z
    niego (stary format)."""
    j = U.plik(katalog, f"{nazwa}_ustawienia.json")
    if j.exists():
        return json.loads(j.read_text(encoding="utf-8"))

    t_ = U.plik(katalog, "00_PODSUMOWANIE.txt")
    tekst = t_.read_text(encoding="utf-8", errors="replace") \
        if t_.exists() else ""
    if "USTAWIENIA, PRZY KTORYCH" not in tekst:
        # juz przepisany przez ten raport - stare wartosci zachowalem obok
        k = U.plik(katalog, "ustawienia_kroku1_odczytane.json")
        return json.loads(k.read_text(encoding="utf-8")) if k.exists() else {}

    def pole(klucz):
        m = re.search(rf"^\s*{re.escape(klucz)}\s*:\s*(.*)$", tekst, re.M)
        return m.group(1).strip() if m else ""

    def num(s, i=0):
        w = re.findall(r"[-+]?\d+(?:\.\d+)?", s)
        return float(w[i]) if len(w) > i else None

    u = {}
    u["plik"] = pole("plik zrodlowy")
    m = re.search(r"\(([\d.]+) ha\)", pole("powierzchnia poletka"))
    u["powierzchnia_ha"] = float(m.group(1)) if m else None
    m = re.search(r"prostokat opisany mialby (\d+) m2", tekst)
    u["prostokat_ha"] = float(m.group(1)) / 10_000 if m else None
    u["krok_pola"] = 0.25
    u["metoda"] = pole("metoda segmentacji")
    u["detekcja"] = pole("metoda wykrywania")
    s = pole("sledzenie pnia")
    u["SLEDZENIE_DOL"], u["SLEDZENIE_GORA"], u["SLEDZENIE_KROK"] = (
        num(s, 0), abs(num(s, 1) or 0), num(s, 2))
    u["SLEDZENIE_MIN_PLASTROW"] = int(num(pole("min. plastrow")) or 0)
    u["MIN_DBH_RAPORTOWANIA"] = (num(pole("prog raportowania DBH")) or 0) / 100
    u["DBH_MIN_POKRYCIE_KATOWE"] = num(pole("min. pokrycie katowe"))
    u["SCIANA_MIN"] = num(pole("filtr sciany pnia"))
    u["MAX_PRZYROST_SREDNICY_CM_NA_M"] = num(pole("srednica ku gorze"))
    s = pole("filtr smuklosci h/d")
    u["SMUKLOSC_MIN"], u["SMUKLOSC_MAX"] = num(s, 0), abs(num(s, 1) or 0)
    m = re.search(r"odrzucil (\d+)", s)
    u["odrzucone_smuklosc"] = int(m.group(1)) if m else 0
    u["DBH_MIN_PUSTOSC"] = num(pole("min. pustosc"))
    u["KORYGUJ_NACHYLENIE"] = pole("korekta nachylenia") == "True"
    U.plik(katalog, "ustawienia_kroku1_odczytane.json").parent.mkdir(
        parents=True, exist_ok=True)
    U.plik(katalog, "ustawienia_kroku1_odczytane.json").write_text(
        json.dumps(u, ensure_ascii=False, indent=1), encoding="utf-8")
    return u


def opisz_ustawienia(u, tab, plik_tab, z_gory=None):
    """Lista (nazwa, wartosc, wyjasnienie) - te same zdania ida do Excela
    i do 00_PODSUMOWANIE.txt."""
    if not u:
        return [(t("Ustawienia kroku 1", "Step 1 settings"),
                 t("brak danych", "no data"),
                 t("Nie znalazłem pliku z ustawieniami wykrywania drzew.",
                   "The tree detection settings file was not found."))]

    def wyl(x):
        return not x

    R = []
    if u.get("metoda") == "graf":
        R.append((t("Podział chmury na drzewa", "Splitting the cloud into "
                    "trees"), t("grafowy", "graph-based"),
                  t("Punkt należy do pnia, do którego prowadzi najkrótsza "
                    "droga przez chmurę punktów (korony nie przeskakują na "
                    "sąsiada).", "A point belongs to the stem reached by "
                    "the shortest path through the point cloud (crowns do "
                    "not jump to the neighbour).")))
    else:
        R.append((t("Podział chmury na drzewa", "Splitting the cloud into "
                    "trees"), t("prosty", "simple"),
                  t("Punkt należy do najbliższego pnia w rzucie poziomym.",
                    "A point belongs to the nearest stem in plan view.")))
    if u.get("detekcja") == "sledzenie":
        R.append((t("Wykrywanie pni", "Stem detection"),
                  t("śledzenie strzały", "stem tracking"),
                  t("Pień jest szukany przekrój po przekroju w górę strzały, "
                    "a nie w jednym plastrze.", "The stem is searched "
                    "section by section up the bole, not in one slice.")))
    else:
        R.append((t("Wykrywanie pni", "Stem detection"),
                  t("przekrój na 1,3 m", "cross-section at 1.3 m"),
                  t("Pnie szukane w jednym poziomym plastrze chmury.",
                    "Stems searched in one horizontal slice of the cloud.")))
    R.append((t("Zakres śledzenia strzały", "Stem tracking range"),
              f"{pl(u['SLEDZENIE_DOL'], 2)}–{pl(u['SLEDZENIE_GORA'], 2)} m, "
              + t(f"co {pl(u['SLEDZENIE_KROK'], 2)} m",
                  f"every {pl(u['SLEDZENIE_KROK'], 2)} m"),
              t("Na jakiej wysokości nad gruntem szukam pnia i co ile "
                "metrów robię przekrój.", "Height range above ground in "
                "which the stem is searched, and the spacing of the "
                "cross-sections.")))
    n = u["SLEDZENIE_MIN_PLASTROW"]
    R.append((t("Minimalna długość strzały", "Minimum stem length"),
              t(f"{n} przekrojów ({pl(n * u['SLEDZENIE_KROK'], 2)} m)",
                f"{n} cross-sections ({pl(n * u['SLEDZENIE_KROK'], 2)} m)"),
              t("Krótsze obiekty (krzewy, pojedyncze gałęzie) nie są "
                "uznawane za pień.", "Shorter objects (shrubs, single "
                "branches) are not accepted as a stem.")))
    R.append((t("Próg pierśnicy", "DBH threshold"),
              f"{pl(u['MIN_DBH_RAPORTOWANIA'] * 100)} cm",
              t("Drzewa cieńsze w pierśnicy nie trafiają do wyników "
                "(jak próg pierśnicy przy inwentaryzacji).", "Trees with a "
                "smaller DBH are not included (like the DBH threshold in a "
                "forest inventory).")))
    R.append((t("Pokrycie obwodu pnia punktami", "Stem circumference "
                "covered by points"),
              t(f"min. {u['DBH_MIN_POKRYCIE_KATOWE']:g}° z 360°",
                f"min. {u['DBH_MIN_POKRYCIE_KATOWE']:g}° of 360°"),
              t("Pierśnica jest mierzona tylko, gdy skaner widział "
                "co najmniej tyle obwodu pnia.", "DBH is measured only when "
                "the scanner saw at least this much of the circumference.")))
    p = u.get("MAX_PRZYROST_SREDNICY_CM_NA_M")
    R.append((t("Kontrola zbieżności strzały", "Stem taper check"),
              t("wyłączona", "off") if wyl(p) else f"maks. +{pl(p)} cm/m"
              if not J.EN else f"max. +{pl(p)} cm/m",
              t("Odrzuca obiekt, którego średnica rośnie ku górze (pień "
                "zwęża się ku wierzchołkowi).", "Rejects an object whose "
                "diameter increases upwards (a stem tapers towards the "
                "top).") if not wyl(p) else
              t("Nie odrzucam pni na podstawie zmiany średnicy z wysokością.",
                "Stems are not rejected because of diameter change with "
                "height.")))
    p = u.get("SCIANA_MIN")
    R.append((t("Test stałości widocznej strony pnia", "Visible stem side "
                "consistency test"),
              t("wyłączony", "off") if wyl(p) else
              t(f"próg {pl(p, 2)}", f"threshold {pl(p, 2)}"),
              t("Prawdziwy pień skaner widzi z tej samej strony na każdej "
                "wysokości; kępa gałęzi – raz z jednej, raz z drugiej.",
                "The scanner sees a real stem from the same side at every "
                "height; a cluster of branches – sometimes from one side, "
                "sometimes from the other.")))
    lo, hi = u.get("SMUKLOSC_MIN"), u.get("SMUKLOSC_MAX")
    if hi:
        w = f"{lo:g}–{hi:g}"
    elif lo:
        w = f"min. {lo:g}"
    else:
        w = t("wyłączony", "off")
    od = u.get("odrzucone_smuklosc", 0)
    R.append((t("Kontrola smukłości h/d", "Slenderness check h/d"), w,
              t("Smukłość = wysokość / pierśnica. Drzewa z wartością poza "
                "zakresem są odrzucane jako błędny pomiar"
                + (f" – odrzucono {od}." if od else "."),
                "Slenderness = height / DBH. Trees outside the range are "
                "rejected as a measurement error"
                + (f" – {od} rejected." if od else "."))))
    R.append((t("Test „pień czy krzak”", "“Stem or shrub” test"),
              t(f"min. {u['DBH_MIN_PUSTOSC'] * 100:.0f}% punktów na obwodzie",
                f"min. {u['DBH_MIN_PUSTOSC'] * 100:.0f}% of points on the "
                f"circumference"),
              t("Przekrój pnia jest pusty w środku (punkty leżą na korze); "
                "krzak ma punkty w całym przekroju.", "A stem cross-section "
                "is hollow (the points lie on the bark); a shrub has points "
                "across the whole section.")))
    R.append((t("Korekta pochylenia pnia", "Stem lean correction"),
              t("włączona", "on") if u.get("KORYGUJ_NACHYLENIE")
              else t("wyłączona", "off"),
              t("Pierśnica liczona prostopadle do osi pochylonego pnia, a nie "
                "w poziomym przekroju (który byłby owalem).", "DBH measured "
                "perpendicular to the axis of a leaning stem, not in the "
                "horizontal section (which would be an oval).")))
    if z_gory:
        if z_gory.get("z_gory"):
            R.append((t("Szukanie pni od góry", "Stem search from above"),
                      f"{pl(z_gory['od'])}–{pl(z_gory['do'])} m, " +
                      t(f"dodano {z_gory.get('dodane', 0)}",
                        f"added {z_gory.get('dodane', 0)}"),
                      t("Drugie szukanie pni wyżej, gdzie strzała jest "
                        "czysta – dla drzew, których dół ginie w podszycie.",
                        "A second stem search higher up, where the bole is "
                        "clear – for trees whose base is hidden in the "
                        "understory.")))
        else:
            R.append((t("Szukanie pni od góry", "Stem search from above"),
                      t("wyłączone", "off"), ""))
    R.append((t("Tablice miąższości", "Volume tables"), plik_tab.name,
              t("Wzory miąższości grubizny dla grup gatunkowych. Plik można "
                "podmienić na tablice regionalne (opis w środku pliku).",
                "Merchantable timber volume formulas (≥7 cm) for species "
                "groups. The file can be replaced with regional tables "
                "(description inside the file).")))
    for gr, w in tab.iterrows():
        if J.EN:
            R.append((f"Volume: {J.gatunek(gr)}", opis_wzoru(w),
                      f"{w.get('opis_en') or w.get('opis', '')}. Source: "
                      f"{w.get('zrodlo_en') or w.get('zrodlo', '')}."))
        else:
            R.append((f"Miąższość: {gr}", opis_wzoru(w),
                      f"{w.get('opis', '')}. Źródło: {w.get('zrodlo', '')}."))
    R.append((t("Dorysowanie zasłoniętego dołu pnia", "Reconstruction of a "
                "hidden stem base"), t("włączone", "on"),
              t("Gdy dół pnia zasłania podszyt, LidarTreeSegmenter szuka czystej strzały "
                "wyżej i przedłuża ją do gruntu (na obrazkach na niebiesko).",
                f"When the stem base is hidden by understory, {J.NAZWA} "
                f"looks for a clear bole higher up and extends it to the "
                f"ground (blue in the images).")))
    return R


def uwaga_pl(kod):
    """Kody uwag z kroku 1 na zwykłe zdania."""
    slownik = {
        "ok": "",
        "niepewny_czesciowy_luk": t("niepewna – widoczna tylko część obwodu",
                                    "uncertain – only part of the "
                                    "circumference visible"),
        "nie_zmierzono_na_1.3m_patrz_srednica_zmierzona":
            t("brak pomiaru na 1,3 m (pień zasłonięty lub to krzew)",
              "not measured at 1.3 m (stem hidden or it is a shrub)"),
    }
    kod = str(kod)
    return slownik.get(kod, kod.replace("_", " "))


# ----------------------------------------------------------------------
# miąższość
# ----------------------------------------------------------------------
def wczytaj_tablice(plik):
    tab = pd.read_csv(plik, sep=";", comment="#", encoding="utf-8-sig",
                      dtype={"grupa": str})
    tab["grupa"] = tab.grupa.str.strip()
    if J.EN:                       # wiersze dopisane po angielsku (larch)
        tab["grupa"] = tab.grupa.map(J.gatunek_wew)
    return tab.set_index("grupa")


def opis_wzoru(w):
    if str(w.wzor).strip() == "potegowy":
        return f"V = {w.a:g} · d^{w.b:g} · h^{w.c:g}"
    return f"V = g · h · {pl(float(w.f), 2)}"


def miazszosc_grubizny(d_m, h, grupa, tab):
    """V [m³] dla każdego drzewa wg wzoru jego grupy gatunkowej."""
    v = np.full(len(d_m), np.nan)
    for i, (d, hh, gr) in enumerate(zip(d_m, h, grupa)):
        w = tab.loc[wiersz_tablic(gr, tab)]
        if not (np.isfinite(d) and np.isfinite(hh)):
            continue
        if str(w.wzor).strip() == "potegowy":
            v[i] = float(w.a) * (d * 100) ** float(w.b) * hh ** float(w.c)
        else:
            v[i] = np.pi / 4 * d ** 2 * hh * float(w.f)
    return v


def wczytaj_sprawdzenie(kat, nazwa):
    """Decyzje z widoku 3D: {nr: (decyzja, gatunek)}. Plik zapisuje
    przycisk „Zapisz sprawdzenie” w widoku 3D: <skan>_sprawdzenie.csv
    (nr;decyzja;gatunek) albo w wersji angielskiej <skan>_review.csv
    (nr;decision;species, wartości tree/shrub/snag, spruce/beech).
    Gdy są oba, nowszy ma pierwszeństwo."""
    wyn = {}
    for p in U.pliki(kat, "sprawdzenie", nazwa):
        try:
            r = pd.read_csv(p, sep=";", encoding="utf-8-sig", dtype=str)
        except (OSError, ValueError):
            continue
        for _, w in r.iterrows():
            try:
                nr_ = int(float(w.get("nr")))
            except (TypeError, ValueError):
                continue
            dec = J.decyzja_wew(w.get("decyzja", w.get("decision", ""))
                                or "")
            if dec == "nan":
                dec = ""
            gat = str(w.get("gatunek", w.get("species", "")) or "") \
                .strip().lower()
            gat = "" if gat in ("nan", "-") else J.gatunek_wew(gat)
            wyn[nr_] = (dec, gat)
    return wyn


def grupy_gatunkowe(kat, nazwa, nr, cz, domyslna, cala=None, etykiety=None):
    """Gatunek każdego drzewa: wpisany ręcznie w <skan>_gatunki.csv (jeśli
    jest), potem z etykiety zapisanej skanerem w terenie (markPoint.csv),
    potem wybrany dla całej powierzchni, potem rozpoznany automatycznie w
    kroku 2 (iglaste = świerk, liściaste = buk), na końcu domyślny.
    Szablon pliku tworzę sam."""
    gr = pd.Series(GATUNEK_GRUPY.get(domyslna, domyslna), index=nr)
    zrodlo = pd.Series(t("domyślny", "default"), index=nr)
    if not cz.empty and "grupa_gatunkowa" in cz:
        a = cz.set_index("treeID").grupa_gatunkowa.dropna()
        a = a[a != "nieokreślone"]
        a = a[a.index.isin(nr)]
        gr[a.index] = a.map(lambda g: GATUNEK_GRUPY.get(g, g))
        zrodlo[a.index] = t("rozpoznany (wstępnie)", "detected "
                            "(preliminary)")
    if cala:                                  # wybrany dla całej powierzchni
        gr[:] = cala
        zrodlo[:] = t("cała powierzchnia", "whole plot")
    for nr_, e in (etykiety or {}).items():       # etykiety ze skanera
        if nr_ in gr.index:
            gr[nr_] = e["gatunek"]
            zrodlo[nr_] = t(f"etykieta skanera ({e['kod']})",
                            f"scanner label ({e['kod']})")
    # <skan>_gatunki.csv, w wersji angielskiej <skan>_species.csv
    plik = U.znajdz(kat, "gatunki", nazwa)
    if plik is not None:
        r = pd.read_csv(plik, sep=";", encoding="utf-8-sig", dtype=str)
        kol = next((k for k in ("gatunek_reczny", "grupa_reczna",
                                "species_manual") if k in r), None)
        if kol:
            r = r.dropna(subset=[kol])
            r[kol] = r[kol].str.strip().str.lower().map(J.gatunek_wew)
            r = r[r[kol] != ""]
            r["nr"] = pd.to_numeric(r.nr, errors="coerce")
            r = r[r.nr.isin(nr)]
            gr[r.nr.astype(int).to_numpy()] = \
                r[kol].map(lambda g: GATUNEK_GRUPY.get(g, g)).to_numpy()
            zrodlo[r.nr.astype(int).to_numpy()] = t("wpisany ręcznie",
                                                    "entered manually")
    else:
        plik = U.plik(kat, U.nazwa_pliku("gatunki", nazwa))
        plik.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame({"nr": nr, t("gatunek_rozpoznany", "species_detected"):
                      [J.gatunek(g) for g in gr.to_numpy()],
                      t("gatunek_reczny", "species_manual"): ""}).to_csv(
            plik, sep=";", index=False, encoding="utf-8-sig")
    return gr.to_numpy(), zrodlo.to_numpy()


def miazszosc_zmierzona(okregi):
    """Suma walców ściętych między kolejnymi okręgami pnia."""
    wynik = {}
    if okregi is None or okregi.empty:
        return wynik
    kol = ("srednica_wygladzona_m" if "srednica_wygladzona_m" in okregi
           else "srednica_m")
    for t, g in okregi.groupby("treeID"):
        g = g.sort_values("wysokosc_m")
        h = g.wysokosc_m.to_numpy()
        d = g[kol].to_numpy()
        if len(h) < 2:
            continue
        dh = np.diff(h)
        ok = dh <= 0.61          # nie przez dziury dłuższe niż 3 przekroje
        a = np.pi / 4 * d ** 2
        v = dh * (a[:-1] + a[1:] + np.sqrt(a[:-1] * a[1:])) / 3
        wynik[int(t)] = float(v[ok].sum())
    return wynik


# ----------------------------------------------------------------------
# Excel - pomocnicze
# ----------------------------------------------------------------------
CIENKA = Side(style="thin", color="BFC9D4")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)


def naglowek(ws, wiersz, kolumny, wysokosc=32):
    for i, tekst in enumerate(kolumny, start=1):
        c = ws.cell(row=wiersz, column=i, value=tekst)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=GRANAT)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = RAMKA
    ws.row_dimensions[wiersz].height = wysokosc


def tytul(ws, tekst, podtytul=None):
    ws["A1"] = tekst
    ws["A1"].font = Font(bold=True, size=16, color=GRANAT)
    if podtytul:
        ws["A2"] = podtytul
        ws["A2"].font = Font(italic=True, color=SZARY)


def krzywa_wysokosci(d, dx, h, uwagi, wolno):
    """Krzywa wysokości drzewostanu (Näslund: h = 1,3 + d² / (a + b·d)²)
    z drzew, na które obie metody dają tę samą pierśnicę (±15%). Gdy krok 1
    i LidarTreeSegmenter różnią się o ponad 30%, bierzemy pomiar, który lepiej pasuje do
    wysokości drzewa (co najmniej o 2 m). Drzewa, których pierśnica wcale
    nie pasuje do wysokości, dostają uwagę „sprawdź”."""
    from scipy.optimize import curve_fit
    d, uwagi = d.copy(), uwagi.copy()
    nas = lambda x, a, b: 1.3 + x ** 2 / (a + b * x) ** 2   # noqa: E731
    with np.errstate(divide="ignore", invalid="ignore"):
        zgodne = (np.isfinite(d) & np.isfinite(dx) & (h > 3) & wolno
                  & (np.abs(d - dx) / np.fmin(d, dx) < 0.15))
        if zgodne.sum() < 8:
            return d, uwagi
        try:
            p, _ = curve_fit(nas, d[zgodne] * 100, h[zgodne], p0=[2, 0.2],
                             maxfev=20000)
        except Exception:                            # noqa: BLE001
            return d, uwagi
        rmse = float(np.sqrt(np.mean((h[zgodne] - nas(d[zgodne] * 100, *p))
                                     ** 2)))
        e1 = np.abs(h - nas(d * 100, *p))
        ex = np.abs(h - nas(dx * 100, *p))
        rozne = (np.isfinite(d) & np.isfinite(dx) & wolno
                 & (np.abs(d - dx) / np.fmin(d, dx) > 0.3))
        zamien = rozne & (ex < e1 - 2.0)
        dziwne = (np.isfinite(d) & np.isfinite(h) & wolno & ~zamien
                  & (np.minimum(e1, np.where(rozne, ex, np.inf))
                     > max(6.0, 2.5 * rmse)))
    for i in np.flatnonzero(zamien):
        uwagi[i] = t(f"pierśnica z LidarTreeSegmenter – lepiej pasuje do wysokości "
                     f"(krok 1 dawał {d[i] * 100:.0f} cm)",
                     f"DBH from {J.NAZWA} – fits the height better "
                     f"(step 1 gave {d[i] * 100:.0f} cm)")
    for i in np.flatnonzero(dziwne):
        if not uwagi[i]:
            uwagi[i] = t("sprawdź na obrazku: pierśnica nie pasuje do "
                         "wysokości drzewa", "check the image: DBH does "
                         "not fit the tree height")
    d = np.where(zamien, dx, d)
    print(t(f"[RAPORT] Krzywa wysokości z {zgodne.sum()} drzew (błąd "
            f"{rmse:.1f} m); pierśnica z LidarTreeSegmenter dla {zamien.sum()} drzew",
            f"[REPORT] Height curve from {zgodne.sum()} trees (error "
            f"{rmse:.1f} m); DBH from {J.NAZWA} for {zamien.sum()} trees"),
          flush=True)
    return d, uwagi


# ----------------------------------------------------------------------
def kontrola_walca(kat, nazwa, drz, cz, okr, u):
    """Kontrola „czy w miejscu drzewa stoi pień-walec” (lts_walec.py).
    Wynik zapisany w <nazwa>_kontrola_walca.csv; liczony od nowa tylko,
    gdy chmura z kroku 2 jest nowsza (Przelicz raport go nie powtarza)."""
    p_w = U.plik(kat, f"{nazwa}_kontrola_walca.csv")
    p_cz = U.plik(kat, f"{nazwa}_czesci.laz")
    p_sg = U.plik(kat, f"{nazwa}_segmenty.laz")
    if not (p_cz.exists() and p_sg.exists()):
        return None
    if p_w.exists() and p_w.stat().st_mtime >= max(p_cz.stat().st_mtime,
                                                   p_sg.stat().st_mtime):
        W = pd.read_csv(p_w)
        if set(drz.id_drzewa.astype(int)) <= set(W.nr):
            return W
    import laspy
    print(t("[RAPORT] Kontrola pnia: czy w miejscu każdego drzewa stoi "
            "walec (prosta oś, stały promień, pusty środek, kora z każdej "
            "strony widzianej ze ścieżki)",
            "[REPORT] Stem check: is there a cylinder at every tree "
            "(straight axis, constant radius, hollow inside, bark from "
            "every side seen from the path)"), flush=True)
    lc, ls = laspy.read(str(p_cz)), laspy.read(str(p_sg))
    if len(lc.points) != len(ls.points):
        raise ValueError("czesci.laz i segmenty.laz mają różną liczbę "
                         "punktów")
    pole_h = next((p for p in LW.POLE_WYSOKOSCI
                   if p in {d_.name for d_ in ls.point_format.dimensions}),
                  None)
    if pole_h is None:
        raise ValueError("brak wysokości nad gruntem w segmenty.laz")
    hag = np.asarray(ls[pole_h], np.float32)
    m = (hag > LW.PAS_OD - 0.1) & (hag < LW.PAS_DO + 0.1)
    xyz = np.column_stack([np.asarray(lc.x)[m], np.asarray(lc.y)[m],
                           np.asarray(lc.z)[m]])
    drzewo = np.asarray(lc.treeID)[m].astype(np.int64)
    czesc = np.asarray(lc.czesc)[m] if "czesc" in \
        {d_.name for d_ in lc.point_format.dimensions} else None
    hag = hag[m]
    del lc, ls
    # kandydaci na środek pnia: krok 1, LidarTreeSegmenter, przekrój pnia na 1,3 m
    D = pd.DataFrame({"nr": drz.id_drzewa.astype(int),
                      "x": drz.pozycja_x.astype(float),
                      "y": drz.pozycja_y.astype(float)})
    if not cz.empty and {"treeID", "pozycja_x", "pozycja_y"} <= set(cz):
        c_ = cz.set_index("treeID")
        D["x2"] = D.nr.map(c_.pozycja_x).astype(float)
        D["y2"] = D.nr.map(c_.pozycja_y).astype(float)
    if len(okr) and {"treeID", "wysokosc_m", "srodek_x"} <= set(okr):
        o = okr[okr.get("zrodlo", "pomiar") == "pomiar"].copy()
        o["dh"] = (o.wysokosc_m - 1.3).abs()
        o = o.sort_values("dh").groupby("treeID").first()
        o = o[o.dh <= 0.5]
        D["x3"] = D.nr.map(o.srodek_x).astype(float)
        D["y3"] = D.nr.map(o.srodek_y).astype(float)
    sektory = None
    sc = SC.dla_skanu(kat, nazwa, u.get("plik"), cicho=True,
                      xy_skanu=D[["x", "y"]].to_numpy(float))
    if sc:
        _, _, sekt = SC.widocznosc(D[["x", "y"]].to_numpy(float), sc["xyz"])
        sektory = dict(zip(D.nr, sekt))
    W = LW.kontrola_skanu(xyz, hag, drzewo, czesc, D, sektory,
                          wypisz=lambda s_: print(s_, flush=True))
    W.to_csv(p_w, index=False)
    return W


def oznacz_odrzucone(kat, numery, odrzucone, powody):
    """Obrazki drzew pominiętych w wynikach (krzewy, brak pnia, duplikaty)
    trafiają do 04_OBRAZY_DRZEW\\POMINIETE z czerwonym napisem i powodem.
    Oryginały leżą w POMINIETE\\_oryginaly, więc po zmianie decyzji
    (np. sprawdzenie drzewa) obrazek wraca na miejsce."""
    import shutil
    kat_o = Path(kat) / U.OBRAZY
    if not kat_o.is_dir():
        return
    pom = kat_o / t("POMINIETE", "EXCLUDED")
    org = pom / "_oryginaly"
    if org.is_dir():                      # najpierw wszystko z powrotem
        for p in org.glob("drzewo_*.png"):
            if (kat_o / p.name).exists():  # nowszy obrazek z kroku 2
                p.unlink()
            else:
                shutil.move(str(p), str(kat_o / p.name))
        for p in pom.glob("drzewo_*.png"):
            p.unlink()
    pliki = {int(re.sub(r"\D", "", p.stem) or -1): p
             for p in kat_o.glob("drzewo_*.png")}
    ile = 0
    for nr_, odr, pw in zip(numery, odrzucone, powody):
        p = pliki.get(int(nr_))
        if not odr or p is None:
            continue
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            org.mkdir(parents=True, exist_ok=True)
            img = plt.imread(str(p))
            hgt, wid = img.shape[:2]
            fig = plt.figure(figsize=(wid / 100, hgt / 100), dpi=100)
            ax = fig.add_axes([0, 0, 1, 1])
            ax.imshow(img)
            ax.axis("off")
            ax.text(0.5, 0.5, t("POMINIĘTE", "EXCLUDED"),
                    transform=ax.transAxes, fontsize=60, color="red",
                    alpha=0.35, ha="center", va="center", rotation=20,
                    weight="bold")
            ax.text(0.01, 0.01, str(pw)[:160], transform=ax.transAxes,
                    fontsize=11, color="white", ha="left", va="bottom",
                    bbox=dict(facecolor="red", alpha=0.85, pad=4))
            fig.savefig(str(pom / p.name), dpi=100)
            plt.close(fig)
            shutil.move(str(p), str(org / p.name))
            ile += 1
        except Exception as e:                     # noqa: BLE001
            print(t(f"  (obrazek {p.name} nieoznaczony: {e})",
                    f"  (image {p.name} not marked: {e})"))
    if ile:
        print(t(f"[RAPORT] Obrazki {ile} pominiętych drzew przeniesione do "
                f"{U.OBRAZY}\\{pom.name}",
                f"[REPORT] Images of {ile} excluded trees moved to "
                f"{U.OBRAZY}\\{pom.name}"), flush=True)


def main():
    ap = argparse.ArgumentParser(description=t(
        __doc__, f"{J.NAZWA} - step 3: REPORT (Excel + text summary)"),
                                 formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("katalog", type=Path,
                    help=t("folder wyników jednego skanu (LidarTreeSegmenter\\wyniki\\<nazwa>)",
                           "results folder of one scan (LidarTreeSegmenter\\wyniki\\<name>)"))
    ap.add_argument("--tablice", type=Path, default=TABLICE,
                    help=t("plik z wzorami miąższości (domyślnie "
                           "tablice_miazszosci.csv obok programu)",
                           "file with volume formulas (default "
                           "tablice_miazszosci.csv next to the program)"))
    ap.add_argument("--grupa", "--group", dest="grupa", default=GRUPA_DOMYSLNA,
                    help=t("grupa gatunkowa dla drzew bez rozpoznanej/wpisanej "
                           "(np. iglaste, liściaste)", "species group for "
                           "trees without a detected/entered one (e.g. "
                           "coniferous, broadleaved)"))
    ap.add_argument("--gatunek", "--species", dest="gatunek", default=None,
                    help=t("jeden gatunek dla całej powierzchni (np. świerk), "
                           "auto = rozpoznawanie z kształtu", "one species "
                           "for the whole plot (e.g. spruce), auto = "
                           "detection from tree shape"))
    ap.add_argument("--epsg", type=int, default=None,
                    help=t("układ dla dopasowania do pomiaru terenowego "
                           "(domyślnie strefa UTM; 2180 = PUWG 1992)",
                           "coordinate system for matching the field "
                           "measurement (default UTM zone; 2180 = PUWG 1992)"))
    args = ap.parse_args()
    args.grupa = J.gatunek_wew(args.grupa)
    if args.gatunek:
        args.gatunek = J.gatunek_wew(args.gatunek)

    kat = args.katalog.resolve()
    if kat.is_file():
        kat = kat.parent
    nazwa = kat.name
    if kat.name in U.PODFOLDERY:           # podano podfolder zamiast skanu
        kat = kat.parent
        nazwa = kat.name
    U.uporzadkuj(kat)
    tab = wczytaj_tablice(args.tablice)
    print(t(f"[RAPORT] Zbieram wyniki: {kat}",
            f"[REPORT] Collecting results: {kat}"), flush=True)

    p_drz = U.plik(kat, f"{nazwa}_drzewa.csv")
    if not p_drz.exists():
        print(t(f"BLAD: brak {p_drz.name} - najpierw krok 1 (wykrywanie "
                f"drzew)", f"ERROR: {p_drz.name} is missing - run step 1 "
                f"(tree detection) first"))
        sys.exit(1)
    drz = pd.read_csv(p_drz)
    p_cz = U.plik(kat, f"{nazwa}_czesci_drzewa.csv")
    cz = pd.read_csv(p_cz) if p_cz.exists() else pd.DataFrame()
    # drzewa, które zna tylko krok 2 (LidarTreeSegmenter śledzi każdy segment, krok 1
    # pomija te bez pomiaru na 1,3 m) - też do raportu, żeby złomy i krzewy
    # nie znikały po cichu
    if not cz.empty and {"treeID", "pozycja_x", "pozycja_y"} <= set(cz):
        brak = cz[~cz.treeID.isin(drz.id_drzewa)]
        if len(brak):
            drz = pd.concat([drz, pd.DataFrame({
                "id_drzewa": brak.treeID.astype(int),
                "pozycja_x": brak.pozycja_x, "pozycja_y": brak.pozycja_y,
                "wysokosc_m": brak.get("wysokosc_drzewa_m", np.nan),
                "dbh_uwagi": t("brak pomiaru w kroku 1 (pierśnica z LidarTreeSegmenter)",
                               f"not measured in step 1 (DBH from "
                               f"{J.NAZWA})"),
            })], ignore_index=True).sort_values("id_drzewa") \
                .reset_index(drop=True)
            for k in ("dbh_m", "dbh_do_uzycia_m"):
                if k in drz:
                    drz[k] = drz[k].astype(float)
            print(t(f"[RAPORT] Drzewa tylko z kroku 2 (LidarTreeSegmenter): {len(brak)} ",
                    f"[REPORT] Trees only from step 2 ({J.NAZWA}): "
                    f"{len(brak)} ") +
                  f"({', '.join(map(str, brak.treeID.astype(int)))})",
                  flush=True)
    p_ok = U.plik(kat, f"{nazwa}_pien_okregi.csv")
    okr = pd.read_csv(p_ok) if p_ok.exists() else pd.DataFrame()
    u = wczytaj_ustawienia(kat, nazwa)

    # --- tabela drzew --------------------------------------------------
    kol_d = "dbh_do_uzycia_m" if "dbh_do_uzycia_m" in drz else "dbh_m"
    d = drz[kol_d].to_numpy(float)
    h = drz.wysokosc_m.to_numpy(float)
    uwagi = drz.get("dbh_uwagi", pd.Series([""] * len(drz))).fillna("") \
        .map(uwaga_pl).to_numpy(object)
    # Brak pierśnicy z kroku 1 (pień zasłonięty na 1,3 m) - bierzemy pomiar
    # LidarTreeSegmenter ze śledzenia całej strzały, jeśli jest wiarygodny: drzewo ma
    # co najmniej 8 m i rozsądną smukłość (40-150). Krzewy zostają puste.
    # kontrola pnia z kroku 2 (LidarTreeSegmenter): krzewy nie wchodzą do wyników
    krzew = np.zeros(len(d), bool)
    kontrola = np.full(len(d), "", object)
    pewna = np.ones(len(d), bool)
    if not cz.empty and "kontrola_pnia" in cz:
        c_ = cz.set_index("treeID")
        kontrola = drz.id_drzewa.map(c_.kontrola_pnia).fillna("") \
            .to_numpy(object)
        krzew = kontrola == "krzew"
        pewna = drz.id_drzewa.map(c_.srednica_13_pewna).fillna(False) \
            .astype(bool).to_numpy()
        powod = drz.id_drzewa.map(c_.kontrola_powod).fillna("") \
            .to_numpy(object) if "kontrola_powod" in c_ \
            else np.full(len(d), "", object)
        uwagi = np.where(krzew, np.array(
            [t("pominięty jako krzew", "excluded as a shrub")
             + (f": {p_}" if p_ else "") for p_ in powod], dtype=object),
                         np.where(kontrola == "niepewny",
                                  t("niepewny: punkty w środku przekroju",
                                    "uncertain: points inside the "
                                    "cross-section"),
                                  np.where(kontrola == "złom", np.array(
                                      [t(f"{p_}; miąższość ze zmierzonej "
                                         f"strzały", f"{p_}; volume from the "
                                         f"measured stem") for p_ in powod],
                                      dtype=object), uwagi)))
    # decyzje ze sprawdzania drzew w widoku 3D mają pierwszeństwo
    spr = wczytaj_sprawdzenie(kat, nazwa)
    krzew, pewna = np.array(krzew, bool), np.array(pewna, bool)
    kontrola, uwagi = np.array(kontrola, object), np.array(uwagi, object)
    # pnie zrośnięte: krok 2 rozdzielił drzewa z dwoma pniami
    p_rz = U.plik(kat, f"{nazwa}_rozdzielone.csv")
    rozdz = pd.read_csv(p_rz) if p_rz.exists() else pd.DataFrame()
    if len(rozdz):
        nr_ix = {int(n): i for i, n in enumerate(drz.id_drzewa.astype(int))}

        def dopisz(n, tekst):
            i = nr_ix.get(int(n))
            if i is not None:
                uwagi[i] = "; ".join(x for x in (str(uwagi[i] or ""), tekst)
                                     if x and x != "nan")
        for _, w in rozdz.iterrows():
            if w.rodzaj == "rozdzielone" and pd.notna(w.nowe):
                dopisz(w.drzewo, t(f"zrośnięte z drzewem {int(w.nowe)} "
                                   f"(rozdzielone)", f"grown together with "
                                   f"tree {int(w.nowe)} (split)"))
                dopisz(w.nowe, t(f"drugi pień drzewa {int(w.drzewo)} "
                                 f"(rozdzielone)", f"second stem of tree "
                                 f"{int(w.drzewo)} (split)"))
            elif w.rodzaj == "rozwidlenie":
                dopisz(w.drzewo, t(f"rozwidlenie na {pl(w.od_m, 1)} m – "
                                   f"jedno drzewo", f"fork at {w.od_m:.1f} m "
                                   f"– one tree"))
    sprawdzone = np.full(len(d), "", object)
    gat_spr = {}
    for i, nr_ in enumerate(drz.id_drzewa.astype(int)):
        if nr_ not in spr:
            continue
        dec, gat_ = spr[nr_]
        if dec == "drzewo":
            krzew[i] = False
            pewna[i] = True
            kontrola[i] = "ok"
            uwagi[i] = t("sprawdzone: drzewo", "reviewed: tree")
        elif dec == "krzew":
            krzew[i] = True
            kontrola[i] = "krzew"
            uwagi[i] = t("sprawdzone: krzew, pominięty w podsumowaniu",
                         "reviewed: shrub, excluded from the summary")
        elif dec == "złom":
            krzew[i] = False
            kontrola[i] = "złom"
            uwagi[i] = t("sprawdzone: złom (złamane drzewo), miąższość ze "
                         "zmierzonej strzały", "reviewed: snag (broken "
                         "tree), volume from the measured stem")
        if dec in ("drzewo", "krzew", "złom"):
            sprawdzone[i] = J.TAK
        if gat_:
            gat_spr[nr_] = gat_
            sprawdzone[i] = J.TAK
    if spr:
        print(t(f"[RAPORT] Uwzględniam sprawdzenie drzew: {len(spr)} decyzji",
                f"[REPORT] Using the tree review: {len(spr)} decisions"),
              flush=True)
    # dwie metody (krok 1 i LidarTreeSegmenter) różnią się o ponad 30% - do sprawdzenia
    if not cz.empty and "srednica_13_m" in cz:
        dx_ = drz.id_drzewa.map(cz.set_index("treeID").srednica_13_m) \
            .to_numpy(float)
        with np.errstate(divide="ignore", invalid="ignore"):
            rozne = (np.abs(dx_ - d) / np.fmin(dx_, d) > 0.3) & ~krzew
        tekst = np.array([t(f"sprawdź na obrazku: pierśnica LidarTreeSegmenter to "
                            f"{x * 100:.0f} cm", f"check the image: "
                            f"{J.NAZWA} DBH is {x * 100:.0f} cm")
                          if np.isfinite(x) else ""
                          for x in dx_], dtype=object)
        uwagi = np.where(rozne & (uwagi == ""), tekst, uwagi)
    uzup = np.zeros(len(d), bool)
    if not cz.empty and "srednica_13_m" in cz:
        dx = drz.id_drzewa.map(cz.set_index("treeID").srednica_13_m) \
            .to_numpy(float)
        with np.errstate(divide="ignore", invalid="ignore"):
            hd_x = h / dx
        uzup = (~np.isfinite(d) & np.isfinite(dx) & (dx > 0) & (h >= 8)
                & (hd_x >= 40) & (hd_x <= 150) & pewna & ~krzew)
        # drugi pień rozdzielonego drzewa: łuk kory od strony pnia głównego
        # jest zasłonięty, a cienki pień w zwarciu bywa bardzo smukły
        drugi = drz.id_drzewa.isin(rozdz.nowe.dropna().astype(int)
                                   if len(rozdz) else []).to_numpy()
        uzup_d = (drugi & ~np.isfinite(d) & np.isfinite(dx) & (dx > 0)
                  & (h >= 8) & (hd_x >= 40) & (hd_x <= 220) & ~krzew)
        d = np.where(uzup, dx, d)
        # złom: pierśnica ze zmierzonego pnia LidarTreeSegmenter, gdy krok 1 jej nie ma
        zl_ = (kontrola == "złom") & ~np.isfinite(d) & np.isfinite(dx)
        d = np.where(zl_, dx, d)
        uwagi = np.where(uzup, t("pierśnica z LidarTreeSegmenter (pień zasłonięty na "
                                 "1,3 m w kroku 1)", f"DBH from {J.NAZWA} "
                                 f"(stem hidden at 1.3 m in step 1)"),
                         uwagi)
        d = np.where(uzup_d, dx, d)
        uwagi = np.where(uzup_d, [t(f"pierśnica z LidarTreeSegmenter – szacunkowa, "
                                    f"drugi pień drzewa {n:.0f}",
                                    f"DBH from {J.NAZWA} – estimate, second "
                                    f"stem of tree {n:.0f}") for n in
                                  drz.id_drzewa.map(dict(zip(
                                      rozdz.nowe.dropna().astype(int),
                                      rozdz.drzewo[rozdz.nowe.notna()]
                                      .astype(int))))
                                  ] if len(rozdz) else uwagi, uwagi)
        uzup = uzup | uzup_d
        d, uwagi = krzywa_wysokosci(d, dx, h, uwagi,
                                    ~krzew & (kontrola != "złom"))
    # kontrola pnia „czy tu stoi walec” (lts_walec.py): obiekty bez
    # pnia i drugi raz policzone pnie nie wchodzą do wyników; decyzje ze
    # sprawdzania drzew mają pierwszeństwo
    odrzut = np.zeros(len(d), bool)
    try:
        KW = kontrola_walca(kat, nazwa, drz, cz, okr, u)
    except Exception as e:                         # noqa: BLE001
        KW = None
        print(t(f"  (kontrola walca pominięta: {e})",
                f"  (cylinder check skipped: {e})"), flush=True)
    if KW is not None and len(KW):
        kw = KW.set_index("nr")
        for i, nr_ in enumerate(drz.id_drzewa.astype(int)):
            if nr_ not in kw.index or (nr_ in spr and spr[nr_][0] in
                                       ("drzewo", "krzew", "złom")):
                continue
            w = kw.loc[nr_]
            if w.werdykt in ("brak_pnia", "duplikat") and not krzew[i]:
                odrzut[i] = True
                krzew[i] = True
                kontrola[i] = "brak pnia" if w.werdykt == "brak_pnia" \
                    else "duplikat"
                uwagi[i] = (t("odrzucony – nie ma tu pnia (walca): ",
                              "rejected – no stem (cylinder) here: ")
                            if w.werdykt == "brak_pnia" else
                            t("odrzucony – ", "rejected – ")) + str(w.powod)
            elif w.werdykt == "pien" and not krzew[i] and \
                    np.isfinite(w.d_walca_m) and \
                    w.pokrycie_deg >= LW.POPRAW_D_POKRYCIE and (
                        not np.isfinite(d[i]) or abs(d[i] - w.d_walca_m)
                        / w.d_walca_m > LW.POPRAW_D_ROZNICA):
                stara = d[i]
                d[i] = w.d_walca_m
                uwagi[i] = t(f"pierśnica z walca pnia ({w.d_walca_m*100:.1f}"
                             f" cm, kora widoczna na {w.pokrycie_deg:.0f}°)",
                             f"DBH from the stem cylinder "
                             f"({w.d_walca_m*100:.1f} cm, bark visible on "
                             f"{w.pokrycie_deg:.0f}°)") + (
                    t(f", wcześniej {stara*100:.1f} cm",
                      f", previously {stara*100:.1f} cm")
                    if np.isfinite(stara) else "")
        print(t(f"[RAPORT] Kontrola walca: odrzucone {int(odrzut.sum())} "
                f"(brak pnia albo pień policzony drugi raz)",
                f"[REPORT] Cylinder check: rejected {int(odrzut.sum())} "
                f"(no stem or stem counted twice)"), flush=True)
    d = np.where(krzew, np.nan, d)            # krzew nie ma pierśnicy
    try:
        oznacz_odrzucone(kat, drz.id_drzewa.astype(int), krzew, uwagi)
    except Exception as e:                         # noqa: BLE001
        print(t(f"  (oznaczanie obrazków pominięte: {e})",
                f"  (marking images skipped: {e})"))
    g = np.pi / 4 * d ** 2
    if args.gatunek:
        UST.zapisz(kat, nazwa, gatunek=args.gatunek)
    cala = UST.gatunek_skanu(kat, nazwa)
    # etykiety gatunku zapisane skanerem w terenie (markPoint.csv)
    et = None
    try:
        et = ET.dla_skanu(kat, nazwa, drz.id_drzewa.astype(int).to_numpy(),
                          drz[["pozycja_x", "pozycja_y"]].to_numpy(float),
                          u.get("plik"))
    except Exception as e:                         # noqa: BLE001
        print(t(f"  (etykiety gatunku pominięte: {e})",
                f"  (species labels skipped: {e})"))
    grupa, grupa_zr = grupy_gatunkowe(kat, nazwa,
                                      drz.id_drzewa.astype(int).to_numpy(),
                                      cz, args.grupa, cala,
                                      et["przypisane"] if et else None)
    for i, nr_ in enumerate(drz.id_drzewa.astype(int)):
        if nr_ in gat_spr:
            grupa[i] = gat_spr[nr_]
            grupa_zr[i] = t("sprawdzony ręcznie", "reviewed manually")
    grupa = np.where(krzew, "nieokreślony", grupa)
    v = miazszosc_grubizny(d, h, grupa, tab)
    for gr in sorted(set(grupa) - set(tab.index)):
        print(t(f"  UWAGA: gatunku „{gr}” nie ma w {args.tablice.name} - "
                f"liczę go wzorem dla „{wiersz_tablic(gr, tab)}” (własny "
                f"wzór dopiszesz w tym pliku)",
                f"  WARNING: species “{J.gatunek(gr)}” is not in "
                f"{args.tablice.name} - using the formula for "
                f"“{J.gatunek(wiersz_tablic(gr, tab))}” (add its own "
                f"formula to that file)"), flush=True)
    vz = miazszosc_zmierzona(okr)
    # złom (złamane drzewo): wzór z wysokością zawyża - bierzemy objętość
    # zmierzonej strzały (walce ścięte), jeśli jest
    zlom = kontrola == "złom"
    if zlom.any():
        vz_ = drz.id_drzewa.map(vz).to_numpy(float)
        v = np.where(zlom & np.isfinite(vz_), vz_, v)

    T = pd.DataFrame({
        "nr": drz.id_drzewa.astype(int),
        "x": drz.pozycja_x.round(2),
        "y": drz.pozycja_y.round(2),
        "d13": (d * 100).round(1),
        "d13_uw": uwagi,
        "h": h.round(1),
        "hd": np.where(d > 0, h / np.where(d > 0, d, 1), np.nan).round(0),
        "g": g.round(4),
        "v": v.round(3),
        "grupa": grupa,
        **({"z_gory": np.where(drz.wykryto_z_gory.fillna(False)
                               .astype(bool), J.TAK, "")}
           if "wykryto_z_gory" in drz else {}),
        "grupa_zr": grupa_zr,
    })
    if not cz.empty:
        c = cz.set_index("treeID")
        def z(k, default=np.nan):
            if k in c:
                return T.nr.map(c[k])
            return pd.Series([default] * len(T), index=T.index)
        T["d13_x"] = (z("srednica_13_m") * 100).round(1)
        T["d13_x_zr"] = z("srednica_13_zrodlo", "").fillna("")
        T["pien_od"] = z("pien_od_m")
        T["pien_do"] = z("pien_do_m")
        T["dor_od"] = z("dorysowano_od_m")
        T["dor_do"] = z("dorysowano_do_m")
        T["proc_pien"] = z("proc_pien")
        T["proc_gal"] = z("proc_galezie")
        T["proc_lis"] = z("proc_ulistnienie")
        T["pkt"] = z("liczba_punktow")
        T["gr_pew"] = z("grupa_pewnosc_proc")
        T["hk"] = z("podstawa_korony_m")
    T["vz"] = T.nr.map(vz).round(3)
    T["kontrola"] = T.nr.map(dict(zip(drz.id_drzewa.astype(int), kontrola)))
    T["sprawdzone"] = T.nr.map(dict(zip(drz.id_drzewa.astype(int),
                                        sprawdzone)))
    T = T.sort_values("nr").reset_index(drop=True)

    # ścieżka skanera (trajektoria), jeśli leży obok skanu
    sc = SC.dla_skanu(kat, nazwa, u.get("plik"),
                      xy_skanu=T[["x", "y"]].to_numpy(float))
    if sc:
        odl_, obej_, _ = SC.widocznosc(T[["x", "y"]].to_numpy(float),
                                       sc["xyz"])
        T["sc_odl"] = np.round(odl_, 1)
        T["sc_obej"] = np.round(obej_, 0)
        T.loc[T.x.isna() | T.y.isna(), ["sc_odl", "sc_obej"]] = np.nan

    # krzewy (podszyt 0,3-4 m) jako osobne obiekty
    try:
        kr, kr_siatka = KR.policz(kat, nazwa)
    except Exception as e:                         # noqa: BLE001
        kr, kr_siatka = None, None
        print(t(f"  (krzewy pominięte: {e})", f"  (shrubs skipped: {e})"))

    # --- drzewostan ----------------------------------------------------
    # krzewy zostają w arkuszu Drzewa (szare), ale nie w podsumowaniu
    licz = ~krzew
    T["krzew"] = T.nr.map(dict(zip(drz.id_drzewa.astype(int), krzew)))
    d, h, g, v, grupa = d[licz], h[licz], g[licz], v[licz], grupa[licz]
    Tl = T[~T.krzew.astype(bool)]
    ha = u.get("powierzchnia_ha") or np.nan
    N = len(Tl)
    zm = np.isfinite(d) & np.isfinite(h)     # drzewa ze zmierzoną pierśnicą
    dg = float(np.sqrt(np.mean(d[zm] ** 2))) if zm.any() else np.nan
    # wysokość Loreya - średnia ważona polem przekroju (standard w taksacji)
    hL = float(np.sum(g[zm] * h[zm]) / np.sum(g[zm])) if zm.any() else np.nan
    g = np.where(np.isfinite(g), g, 0.0)
    v = np.where(np.isfinite(v), v, 0.0)
    POW = t("Powierzchnia próbna", "Sample plot area")
    SZT, SZT_HA = t("szt.", "trees"), t("szt./ha", "trees/ha")
    S = [
        (POW, ha * 10_000, "m²", 0),
        (POW, ha, "ha", 4),
        (t("Liczba drzew", "Number of trees"), N, SZT, 0),
        (t("Liczba drzew na 1 ha (N)", "Trees per hectare (N)"), N / ha,
         SZT_HA, 0),
        (t("Drzewa bez pomiaru pierśnicy", "Trees without DBH measurement"),
         int(N - zm.sum()), SZT, 0),
        (t("Pierśnica uzupełniona z LidarTreeSegmenter", f"DBH filled in from "
           f"{J.NAZWA}"), int(uzup.sum()), SZT, 0),
        (t("Średnia arytmetyczna pierśnica", "Mean DBH (arithmetic)"),
         np.nanmean(d) * 100, "cm", 1),
        (t("Pierśnica przeciętna (Dg, z pola przekroju)",
           "Quadratic mean diameter (Dg)"), dg * 100, "cm", 1),
        (t("Średnia arytmetyczna wysokość", "Mean height (arithmetic)"),
         np.nanmean(h), "m", 1),
        (t("Wysokość Loreya (ważona polem przekroju)",
           "Lorey's height (basal-area weighted)"), hL, "m", 1),
        (t("Pierśnicowe pole przekroju – powierzchnia", "Basal area – "
           "plot"), g.sum(), "m²", 3),
        (t("Pierśnicowe pole przekroju (G)", "Basal area (G)"),
         g.sum() / ha, "m²/ha", 2),
        (t("Miąższość grubizny – powierzchnia", "Merchantable volume "
           "≥7 cm – plot"), v.sum(), "m³", 2),
        (t("Miąższość grubizny (V)", "Merchantable volume ≥7 cm (V)"),
         v.sum() / ha, "m³/ha", 0),
        (t("Średnia miąższość drzewa", "Mean tree volume"),
         v[zm].mean() if zm.any() else np.nan, "m³", 3),
    ]
    if "proc_pien" in T:
        S += [(t("Średni udział punktów: pień", "Mean share of points: "
                 "stem"), Tl.proc_pien.mean(), "%", 1),
              (t("Średni udział punktów: gałęzie", "Mean share of points: "
                 "branches"), Tl.proc_gal.mean(), "%", 1),
              (t("Średni udział punktów: ulistnienie", "Mean share of "
                 "points: foliage"), Tl.proc_lis.mean(), "%", 1)]
    if "d13_x_zr" in T:
        S.append((t("Drzewa z dorysowanym dołem pnia", "Trees with "
                    "reconstructed stem base"),
                  int((Tl.dor_od.notna()).sum()), SZT, 0))
    if (krzew & ~odrzut).any():
        S.append((t("Pominięte jako krzewy (pełny przekrój)", "Excluded as "
                    "shrubs (solid cross-section)"),
                  int((krzew & ~odrzut).sum()), SZT, 0))
    if (kontrola == "brak pnia").any():
        S.append((t("Odrzucone – nie ma pnia (walca)", "Rejected – no stem "
                    "(cylinder)"), int((kontrola == "brak pnia").sum()),
                  SZT, 0))
    if (kontrola == "duplikat").any():
        S.append((t("Odrzucone – pień policzony drugi raz", "Rejected – "
                    "stem counted twice"), int((kontrola == "duplikat")
                                               .sum()), SZT, 0))
    if zlom.any():
        S.append((t("Złomy (złamane drzewa, w wynikach)", "Snags (broken "
                    "trees, included)"), int(zlom.sum()), SZT, 0))

    for gr in sorted(set(grupa)):
        m = grupa == gr
        if len(set(grupa)) > 1:
            S.append((t(f"Miąższość grubizny – {gr}", f"Merchantable volume "
                        f"≥7 cm – {J.gatunek(gr)}"), v[m].sum() / ha,
                      "m³/ha", 0))
    pz = U.plik(kat, f"{nazwa}_szukanie_z_gory.json")
    z_gory = json.loads(pz.read_text(encoding="utf-8")) if pz.exists() \
        else None
    if sc:
        S += [(t("Długość ścieżki skanera", "Scanner path length"),
               sc["dlugosc_m"], "m", 0),
              (t("Średnia odległość drzewa od ścieżki", "Mean distance of a "
                 "tree from the path"), Tl.sc_odl.mean(), "m", 1),
              (t("Średnie obejście drzewa ze ścieżki", "Mean share of the "
                 "tree circumference passed by the path"),
               Tl.sc_obej.mean(), "%", 0),
              (t(f"Drzewa dalej niż {SC.DALEKO_M:.0f} m od ścieżki",
                 f"Trees farther than {SC.DALEKO_M:.0f} m from the path"),
               int((Tl.sc_odl > SC.DALEKO_M).sum()), SZT, 0)]
    if et:
        S.append((t("Drzewa z gatunkiem z etykiety skanera", "Trees with "
                    "species from a scanner label"),
                  int(Tl.nr.isin(list(et["przypisane"])).sum()), SZT, 0))
    # model terenu z kroku 10 (03_TABELE/<n>_teren.npz)
    st_ter = None
    pt = U.plik(kat, f"{nazwa}_teren.npz")
    if pt.exists():
        try:
            st_ter = DTM.statystyki(DTM.wczytaj(pt))
        except Exception as e:                     # noqa: BLE001
            print(t(f"  (model terenu pominięty: {e})",
                    f"  (terrain model skipped: {e})"))
    if st_ter:
        S += [(t("Teren – najniższy punkt", "Terrain – lowest point"),
               st_ter["min"], t("m (układ skanu)", "m (scan system)"), 1),
              (t("Teren – deniwelacja", "Terrain – relief"),
               st_ter["deniwelacja"], "m", 1),
              (t("Teren – średni spadek", "Terrain – mean slope"),
               st_ter["spadek_sr"], "°", 1),
              (t("Teren – średni spadek w procentach",
                 "Terrain – mean slope in percent"),
               st_ter["spadek_proc"], "%", 0),
              (t("Teren – spadek stromych miejsc (95%)",
                 "Terrain – slope of steep places (95%)"),
               st_ter["spadek_max"], "°", 0),
              (t("Teren – ekspozycja stoku", "Terrain – slope aspect"),
               st_ter["ekspozycja_deg"],
               f"° ({st_ter['ekspozycja']}, {st_ter['ekspozycja_opis']})", 0)]
    if len(rozdz):
        S += [(t("Drzewa rozdzielone (dwa pnie w jednym)", "Trees split "
                 "(two stems in one)"),
               int((rozdz.rodzaj == "rozdzielone").sum()), SZT, 0),
              (t("Drzewa z rozwidleniem powyżej 1,3 m", "Trees forked "
                 "above 1.3 m"),
               int((rozdz.rodzaj == "rozwidlenie").sum()), SZT, 0)]
    if kr is not None:
        pole_kr = float(kr.pole_m2.sum()) if len(kr) else 0.0
        S += [(t("Krzewy (0,3–4 m) – liczba", "Shrubs (0.3–4 m) – count"),
               len(kr), SZT if not J.EN else "shrubs", 0),
              (t("Krzewy – na hektar", "Shrubs – per hectare"),
               len(kr) / ha, t("szt./ha", "shrubs/ha"), 0),
              (t("Krzewy – pokrycie powierzchni", "Shrubs – plot cover"),
               pole_kr / (ha * 10_000) * 100, "%", 1),
              (t("Krzewy – średnia wysokość", "Shrubs – mean height"),
               float(kr.wysokosc_m.mean()) if len(kr) else np.nan, "m", 1)]
    if "z_gory" in T:
        S.append((t("Drzewa wykryte dopiero od góry", "Trees detected only "
                    "from above"), int((Tl.z_gory == J.TAK).sum()), SZT, 0))
    opisy = opisz_ustawienia(u, tab, args.tablice, z_gory)
    ux = UST.wczytaj(kat, nazwa)
    kl = ux.get("klasa_uzyta")
    if kl:
        par = UST.klasy().get(kl, {})
        auto_ = J.klasa_wew(ux.get("klasa", "auto")) == "auto"
        opisy.insert(0, (t("Klasa drzewostanu", "Stand class"), J.klasa(kl) + (
            t(" (dobrana automatycznie z pierśnic)",
              " (chosen automatically from DBH)")
            if auto_ else t(" (wybrana)", " (selected)")),
            t(f"{par.get('opis', '')}. Krzew: obiekt bez pustego pnia niższy "
              f"niż {pl(par.get('wysokosc_krzewu_m', 10), 0)} m. Podszyt: "
              f"punkty poniżej {pl(par.get('podszyt_do_m', 0), 1)} m dalej "
              f"niż {pl(par.get('podszyt_odleglosc_m', 0), 1)} m od pnia "
              f"(plik klasy_drzewostanu.csv).",
              f"{UST.opis(kl, par)}. Shrub: an object without a hollow "
              f"stem lower than {pl(par.get('wysokosc_krzewu_m', 10), 0)} m. "
              f"Understory: points below "
              f"{pl(par.get('podszyt_do_m', 0), 1)} m farther than "
              f"{pl(par.get('podszyt_odleglosc_m', 0), 1)} m from the stem "
              f"(file klasy_drzewostanu.csv).")))
    opisy.insert(1 if kl else 0, (t("Gatunek na powierzchni",
                                    "Species on the plot"),
                  J.gatunek(cala) or t("rozpoznawany z kształtu drzewa "
                                       "(wstępnie)", "detected from tree "
                                       "shape (preliminary)"),
                  t("Jeden gatunek dla wszystkich drzew wybierzesz w oknie "
                    "LidarTreeSegmenter albo opcją --gatunek. Wpisy w pliku gatunków mają "
                    "pierwszeństwo.", f"One species for all trees can be "
                    f"chosen in the {J.NAZWA} window or with the --species "
                    f"option. Entries in the species file take precedence.")))
    if et:
        n_et = len(et["wszystkie"])
        n_nie = sum(e["nr"] is None for e in et["wszystkie"])
        opisy.append((t("Etykiety gatunku ze skanera",
                        "Species labels from the scanner"),
                      f"{et['plik'].parent.name}/{et['plik'].name}: "
                      + t(f"{n_et} etykiet, przypisane "
                          f"{len(et['przypisane'])} drzewom",
                          f"{n_et} labels, assigned to "
                          f"{len(et['przypisane'])} trees"),
                      t(f"Etykieta trafia do najbliższego pnia bliżej niż "
                        f"{ET.MAKS_ODL_M:g} m ({n_nie} dalej, bez "
                        f"przypisania; zobacz warstwę etykiety_skanera w "
                        f"{nazwa}.gpkg). Ma pierwszeństwo przed gatunkiem "
                        f"z kształtu i wybranym dla całej powierzchni, "
                        f"ustępuje plikowi gatunków i sprawdzaniu drzew."
                        + (" Nieznane kody: " + ", ".join(
                            et["nieznane_kody"]) + "."
                           if et["nieznane_kody"] else ""),
                        f"A label goes to the nearest stem closer than "
                        f"{ET.MAKS_ODL_M:g} m ({n_nie} farther, not "
                        f"assigned; see the layer scanner_labels in "
                        f"{nazwa}.gpkg). It takes precedence over the "
                        f"species from shape and the one chosen for the "
                        f"whole plot, and yields to the species file and "
                        f"the tree review."
                        + (" Unknown codes: " + ", ".join(
                            et["nieznane_kody"]) + "."
                           if et["nieznane_kody"] else ""))))
    if kr is not None:
        opisy.append((t("Krzewy", "Shrubs"),
                      t(f"{len(kr)} krzewów, arkusz Krzewy",
                        f"{len(kr)} shrubs, sheet Shrubs"),
                      t(f"Krzew to bryła gałęzi i liści 0,3–4 m nad gruntem "
                        f"poza pniami: komórki siatki "
                        f"{pl(KR.KOMORKA_M, 2)} m zajęte w co najmniej "
                        f"{KR.MIN_WARSTW} warstwach 0,5–3,25 m, połączone w "
                        f"płaty i podzielone według wierzchołków, najmniej "
                        f"{pl(KR.MIN_POLE_M2, 1)} m². Liczą się punkty "
                        f"podszytu i niskie gałęzie drzew oddzielone od "
                        f"korony pustym pasem. Pień (i "
                        f"{KR.OD_KORY_M * 100:.0f} cm wokół kory) nigdy nie "
                        f"jest krzewem, więc pierśnica się nie zmienia. W "
                        f"{nazwa}.gpkg warstwa krzewy.",
                        f"A shrub is a body of branches and leaves 0.3–4 m "
                        f"above ground outside the stems: "
                        f"{KR.KOMORKA_M:g} m grid cells filled in at least "
                        f"{KR.MIN_WARSTW} layers between 0.5 and 3.25 m, "
                        f"joined into patches and split by their tops, at "
                        f"least {KR.MIN_POLE_M2:g} m². Understory points and "
                        f"low tree branches separated from the crown by an "
                        f"empty band count. The stem (and "
                        f"{KR.OD_KORY_M * 100:.0f} cm around the bark) is "
                        f"never a shrub, so DBH does not change. Layer "
                        f"shrubs in {nazwa}.gpkg.")))
    if st_ter:
        opisy.append((t("Model terenu", "Terrain model"),
                      t(f"siatka {pl(DTM.KROK_M, 2)} m z punktów gruntu",
                        f"{DTM.KROK_M:g} m grid from ground points"),
                      t("Spadek liczony po wygładzeniu (ok. 0,75 m), "
                        "ekspozycja to kierunek, w który opada stok. W pliku "
                        f"{nazwa}.gpkg są warstwy teren_wysokosc, "
                        "teren_spadek, teren_cieniowanie i warstwice.",
                        "Slope computed after smoothing (about 0.75 m); the "
                        "aspect is the direction the slope faces. The file "
                        f"{nazwa}.gpkg holds the layers terrain_elevation, "
                        "terrain_slope, terrain_hillshade and contours.")))
    if sc:
        opisy.append((t("Ścieżka skanera", "Scanner path"),
                      f"{sc['plik'].name} ({len(sc['xyz']):,} "
                      + t("punktów", "points") + ")",
                      t(f"Odległość od ścieżki i obejście drzewa liczone w "
                        f"poziomie; obejście to część z {SC.SEKTORY} "
                        f"sektorów wokół pnia, w których ścieżka przeszła "
                        f"bliżej niż {SC.ZASIEG_WIDOCZNOSCI_M:.0f} m.",
                        f"Distance from the path and the share of the "
                        f"circumference are horizontal; the share is the "
                        f"part of {SC.SEKTORY} sectors around the stem in "
                        f"which the path passed closer than "
                        f"{SC.ZASIEG_WIDOCZNOSCI_M:.0f} m.")))
    wzory = "; ".join(f"{J.gatunek(gr)}: {opis_wzoru(tab.loc[wiersz_tablic(gr, tab)])}"
                      + ("" if gr in tab.index else
                         f" ({t('jak', 'as')} {J.gatunek(wiersz_tablic(gr, tab))})")
                      for gr in sorted(set(grupa)))

    # --- 00_PODSUMOWANIE.txt ------------------------------------------
    W = 46
    f_gat = U.nazwa_pliku("gatunki", nazwa)
    f_rap = U.nazwa_pliku("raport", nazwa)
    L = ["=" * 90, t(f"{J.NAZWA} – RAPORT SKANU {nazwa}",
                     f"{J.NAZWA} – SCAN REPORT {nazwa}"),
         t(f"utworzono {dt.datetime.now():%Y-%m-%d %H:%M}",
           f"created {dt.datetime.now():%Y-%m-%d %H:%M}"), "=" * 90, "",
         t("PODSUMOWANIE DRZEWOSTANU", "STAND SUMMARY"), "-" * 90,
         f"  {t('Plik źródłowy', 'Source file'):<{W}}: {u.get('plik', '–')}"]
    for nazw, war, jedn, m in S:
        if nazw == POW and jedn == "ha":
            continue
        if nazw == POW:
            L.append(f"  {nazw:<{W}}: {pl(war, 0)} m² ({pl(ha, 4)} ha)")
            continue
        L.append(f"  {nazw:<{W}}: {pl(war, m)} {jedn}")
    L += ["", t(f"  Miąższość grubizny wg {args.tablice.name} ({wzory}).",
                f"  Merchantable timber volume (≥7 cm, over bark) by "
                f"{args.tablice.name} ({wzory})."),
          t(f"  Gatunek drzewa możesz wpisać w {f_gat} "
            f"(kolumna gatunek_reczny) i uruchomić raport jeszcze raz.",
            f"  You can enter the species of each tree in {f_gat} "
            f"(column species_manual) and run the report again."),
          "-" * 90, "",
          t("USTAWIENIA OBLICZEŃ", "CALCULATION SETTINGS"), "-" * 90]
    for nazw, war, wyj in opisy:
        L.append(f"  {nazw:<{W}}: {war}")
        L.append(f"  {'':<{W}}  {wyj}")
    L += ["-" * 90, "",
          t(f"Pełna tabela drzew: {f_rap} (arkusz „Drzewa”).",
            f"Full tree table: {f_rap} (sheet “Trees”).")]
    rap = U.katalog(kat, U.RAPORT)
    (rap / U.nazwa_pliku("podsumowanie")).write_text("\n".join(L) + "\n",
                                                     encoding="utf-8")
    U.usun_inny_jezyk(kat, "podsumowanie")
    print("\n".join(L[5:6 + len(S) + 2]), flush=True)

    # --- Excel ---------------------------------------------------------
    teren = pomiar_terenowy(kat, nazwa, u, T, args.epsg)
    wb = Workbook()

    # Podsumowanie
    ws = wb.active
    ws.title = t("Podsumowanie", "Summary")
    tytul(ws, t(f"Raport drzewostanu – {nazwa}", f"Stand report – {nazwa}"),
          f"{J.NAZWA} · {dt.datetime.now():%Y-%m-%d %H:%M} · "
          f"{u.get('plik', '')}")
    naglowek(ws, 4, [t("Wskaźnik", "Parameter"), t("Wartość", "Value"),
                     t("Jednostka", "Unit")], 22)
    r = 5
    for nazw, war, jedn, m in S:
        ws.cell(row=r, column=1, value=nazw)
        c = ws.cell(row=r, column=2,
                    value=None if not np.isfinite(war) else float(war))
        c.number_format = "0" if m == 0 else "0." + "0" * m
        ws.cell(row=r, column=3, value=jedn)
        wazne = jedn in (SZT_HA, "m²/ha", "m³/ha")
        for k in (1, 2, 3):
            ws.cell(row=r, column=k).border = RAMKA
            if wazne:
                ws.cell(row=r, column=k).font = Font(bold=True)
                ws.cell(row=r, column=k).fill = PatternFill("solid",
                                                            fgColor=JASNY)
        r += 1
    r += 1
    ws.cell(row=r, column=1,
            value=t(f"Miąższość grubizny wg {args.tablice.name}: {wzory}. "
                    f"Grupy gatunkowe wpiszesz w {f_gat}.",
                    f"Merchantable timber volume (≥7 cm, over bark) by "
                    f"{args.tablice.name}: {wzory}. Species can be entered "
                    f"in {f_gat}.")).font = \
        Font(italic=True, color=SZARY)

    # rozkład pierśnic w stopniach co 4 cm
    r += 2
    ws.cell(row=r, column=1, value=t("Rozkład pierśnic",
                                     "DBH distribution")).font = Font(
        bold=True, size=12, color=GRANAT)
    r += 1
    naglowek(ws, r, [t("Stopień grubości [cm]", "DBH class [cm]"),
                     t("Liczba drzew", "Number of trees"),
                     t("Liczba drzew na 1 ha", "Trees per ha"),
                     "G [m²/ha]", "V [m³/ha]"], 30)
    st = np.where(zm, np.floor(np.nan_to_num(d) * 100 / 4) * 4, -1).astype(int)
    for s0 in sorted(set(st) - {-1}):
        r += 1
        m = st == s0
        wiersz = [t(f"{s0}–{s0 + 3.9:.1f}".replace(".", ","),
                    f"{s0}–{s0 + 3.9:.1f}"),
                  int(m.sum()), m.sum() / ha, g[m].sum() / ha, v[m].sum() / ha]
        for k, w in enumerate(wiersz, start=1):
            c = ws.cell(row=r, column=k, value=w)
            c.border = RAMKA
            c.number_format = ["@", "0", "0", "0.00", "0"][k - 1]
    for k, w in zip("ABCDE", (44, 14, 18, 12, 12)):
        ws.column_dimensions[k].width = w

    # Drzewa
    ws = wb.create_sheet(t("Drzewa", "Trees"))
    kol = [  # (klucz, nagłówek, format, szerokość)
        ("nr", t("Nr drzewa", "Tree no."), "0", 8),
        ("x", "X [m]", "0.00", 9),
        ("y", "Y [m]", "0.00", 9),
        ("d13", t("Pierśnica D1,3 [cm]", "DBH D1.3 [cm]"), "0.0", 11),
        ("h", t("Wysokość H [m]", "Tree height H [m]"), "0.0", 10),
        ("g", t("Pole przekroju g [m²]", "Basal area g [m²]"), "0.0000", 11),
        ("v", t("Miąższość grubizny V [m³]", "Merchantable volume ≥7 cm "
                "V [m³]"), "0.000", 12),
        ("hd", t("Smukłość h/d", "Slenderness h/d"), "0", 9),
        ("z_gory", t("Wykryte od góry", "Detected from above"), "@", 9),
        ("sprawdzone", t("Sprawdzone ręcznie", "Reviewed manually"), "@",
         10),
        ("grupa", t("Gatunek", "Species"), "@", 13),
        ("grupa_zr", t("Gatunek – skąd", "Species – source"), "@", 16),
        ("gr_pew", t("Pewność rozpoznania grupy [%]", "Species group "
                     "confidence [%]"), "0", 11),
        ("hk", t("Podstawa korony [m]", "Crown base height [m]"), "0.0", 10),
        ("d13_uw", t("Uwagi do pierśnicy", "DBH notes"), "@", 34),
        ("d13_x", t("Pierśnica LidarTreeSegmenter [cm]", f"DBH {J.NAZWA} [cm]"), "0.0",
         11),
        ("d13_x_zr", t("Pierśnica LidarTreeSegmenter – źródło", f"DBH {J.NAZWA} – "
                       f"source"), "@", 13),
        ("pien_od", t("Strzała widoczna od [m]", "Stem visible from [m]"),
         "0.0", 10),
        ("pien_do", t("Strzała widoczna do [m]", "Stem visible to [m]"),
         "0.0", 10),
        ("dor_od", t("Dół dorysowany od [m]", "Base reconstructed from "
                     "[m]"), "0.0", 10),
        ("dor_do", t("Dół dorysowany do [m]", "Base reconstructed to [m]"),
         "0.0", 10),
        ("vz", t("Miąższość zmierzonej strzały [m³]", "Volume of measured "
                 "stem section [m³]"), "0.000", 12),
        ("proc_pien", t("Pień [% pkt]", "Stem [% pts]"), "0.0", 8),
        ("proc_gal", t("Gałęzie [% pkt]", "Branches [% pts]"), "0.0", 8),
        ("proc_lis", t("Ulistnienie [% pkt]", "Foliage [% pts]"), "0.0", 10),
        ("pkt", t("Liczba punktów", "Number of points"), "#,##0", 11),
        ("sc_odl", t("Odległość od ścieżki skanera [m]", "Distance from the "
                     "scanner path [m]"), "0.0", 11),
        ("sc_obej", t("Obejście ze ścieżki [%]", "Circumference passed by "
                      "the path [%]"), "0", 11),
    ]
    if J.EN:                  # wartości w arkuszu i w tabelach po angielsku
        for k_ in ("grupa", "kontrola", "d13_x_zr"):
            if k_ in T:
                T[k_] = T[k_].map(lambda v: J.WARTOSCI_EN.get(v, v)
                                  if isinstance(v, str) else v)
    kol = [k for k in kol if k[0] in T]
    naglowek(ws, 1, [k[1] for k in kol], 48)
    nieb = PatternFill("solid", fgColor=NIEBIESKI)
    zolt = PatternFill("solid", fgColor=ZOLTY)
    pas = PatternFill("solid", fgColor="F7F9FB")
    Tk = T[T.krzew.astype(bool)].reset_index(drop=True)
    T = T[~T.krzew.astype(bool)].reset_index(drop=True)

    def wiersz_drzewa(r, row, szary=False):
        for j, (k, _, fmt, _) in enumerate(kol, start=1):
            w = row[k]
            if isinstance(w, float) and not np.isfinite(w):
                w = None
            elif isinstance(w, (np.integer,)):
                w = int(w)
            elif isinstance(w, (np.floating,)):
                w = float(w)
            c = ws.cell(row=r, column=j, value=w)
            c.number_format = fmt
            c.border = RAMKA
            if szary:
                c.font = Font(italic=True, color=SZARY)

    for i, row in T.iterrows():
        r = i + 2
        dorysowany = "d13_x_zr" in T and row.get("d13_x_zr") in (
            "rekonstrukcja", "reconstructed")
        uwaga = str(row.get("d13_uw", "")) not in ("", "nan")
        for j, (k, _, fmt, _) in enumerate(kol, start=1):
            w = row[k]
            if isinstance(w, float) and not np.isfinite(w):
                w = None
            elif isinstance(w, (np.integer,)):
                w = int(w)
            elif isinstance(w, (np.floating,)):
                w = float(w)
            c = ws.cell(row=r, column=j, value=w)
            c.number_format = fmt
            c.border = RAMKA
            if i % 2:
                c.fill = pas
            if dorysowany and k in ("d13_x", "d13_x_zr", "dor_od", "dor_do"):
                c.fill = nieb
            if uwaga and k in ("d13", "d13_uw"):
                c.fill = zolt
    # wiersz sumy / średniej (formuły Excela)
    r = len(T) + 2
    ostatni = r - 1
    ws.cell(row=r, column=1, value=t("Razem / średnio", "Total / mean")) \
        .font = Font(bold=True)
    for j, (k, _, fmt, _) in enumerate(kol, start=1):
        lit = get_column_letter(j)
        if k in ("g", "v", "vz", "pkt"):
            f_ = f"=SUM({lit}2:{lit}{ostatni})"
        elif k in ("d13", "h", "hd", "d13_x", "proc_pien", "proc_gal",
                   "proc_lis"):
            f_ = f"=AVERAGE({lit}2:{lit}{ostatni})"
        else:
            continue
        c = ws.cell(row=r, column=j, value=f_)
        c.number_format = fmt
        c.font = Font(bold=True)
    for j in range(1, len(kol) + 1):
        ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor=JASNY)
        ws.cell(row=r, column=j).border = RAMKA
    for j, (_, _, _, szer) in enumerate(kol, start=1):
        ws.column_dimensions[get_column_letter(j)].width = szer
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kol))}{ostatni}"
    r += 2
    for kolor, opis in ((NIEBIESKI, t("pierśnica z dorysowanego "
                                      "(zasłoniętego) dołu pnia – wartość "
                                      "szacunkowa", "DBH from a reconstructed "
                                      "(hidden) stem base – estimated value")),
                        (ZOLTY, t("uwaga przy pomiarze pierśnicy w kroku 1",
                                  "note on the DBH measurement in step 1"))):
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=kolor)
        ws.cell(row=r, column=2, value=opis).font = Font(italic=True,
                                                         color=SZARY)
        r += 1
    if len(Tk):
        r += 1
        ws.cell(row=r, column=1,
                value=t(f"Pominięte ({len(Tk)}) – krzewy (przekrój pełny "
                        f"punktów zamiast pustego pnia), obiekty bez pnia "
                        f"i drugi raz policzone pnie (powód w kolumnie "
                        f"uwag); nie wchodzą do sum i podsumowania:",
                        f"Excluded ({len(Tk)}) – shrubs (cross-section full "
                        f"of points instead of a hollow stem), objects "
                        f"without a stem and stems counted twice (reason "
                        f"in the notes column); not included in totals and "
                        f"the summary:")).font = Font(bold=True, color=SZARY)
        for _, row in Tk.iterrows():
            r += 1
            wiersz_drzewa(r, row, szary=True)

    if kr is not None and len(kr):
        arkusz_krzewow(wb, kr)
    # Ustawienia
    if teren:
        arkusz_terenowy(wb, teren)
    ws = wb.create_sheet(t("Ustawienia", "Settings"))
    tytul(ws, t("Ustawienia obliczeń", "Calculation settings"),
          t("Przy jakich ustawieniach powstały te wyniki",
            "The settings these results were calculated with"))
    naglowek(ws, 4, [t("Ustawienie", "Setting"), t("Wartość", "Value"),
                     t("Co to oznacza", "What it means")], 22)
    for i, (nazw, war, wyj) in enumerate(opisy, start=5):
        for j, w in enumerate((nazw, war, wyj), start=1):
            c = ws.cell(row=i, column=j, value=w)
            c.border = RAMKA
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.row_dimensions[i].height = 32
    for k, w in zip("ABC", (36, 26, 80)):
        ws.column_dimensions[k].width = w

    # Objaśnienia
    ws = wb.create_sheet(t("Objaśnienia", "Explanations"))
    tytul(ws, t("Objaśnienia kolumn arkusza „Drzewa”",
                "Explanation of the columns of the “Trees” sheet"))
    obj = OBJASNIENIA_EN(args.tablice.name, wzory, f_gat) if J.EN else {
        "nr": "Numer drzewa – ten sam co w plikach chmur i na obrazkach "
              "(podglad\\drzewo_NNN.png).",
        "x": "Położenie pnia w układzie współrzędnych skanu.",
        "y": "Położenie pnia w układzie współrzędnych skanu.",
        "d13": "Pierśnica na wysokości 1,3 m, zmierzona w kroku 1 "
               "(wykrywanie drzew). Na niej liczone są G i V.",
        "h": "Wysokość drzewa – najwyższy punkt korony nad gruntem.",
        "g": "Pierśnicowe pole przekroju: π/4 · D1,3².",
        "v": f"Miąższość grubizny wg wzoru grupy gatunkowej z "
             f"{args.tablice.name} ({wzory}).",
        "z_gory": "tak – pień znaleziony dopiero w drugim szukaniu, wyżej "
                  "(dół zasłonięty podszytem).",
        "grupa": "Gatunek, od którego zależy wzór miąższości. Rozpoznana "
                 "grupa iglasta to świerk, liściasta buk (główne gatunki "
                 "na powierzchniach). Inny gatunek wpisz ręcznie.",
        "gr_pew": "Jak wyraźnie kształt drzewa wskazał grupę (reguły: "
                  "pień do wierzchołka, ostry szczyt i stożkowa korona = "
                  "iglaste; rozwidlony pień, szeroka zaokrąglona korona = "
                  "liściaste). Ocena wstępna – sprawdź na obrazku drzewa.",
        "hk": "Wysokość, na której zaczyna się korona.",
        "grupa_zr": "domyślny – przyjęty dla wszystkich; rozpoznany – z "
                    "kształtu drzewa; etykieta skanera – kod zapisany "
                    "skanerem w terenie (markPoint.csv, np. FI = świerk, "
                    "BU = buk); wpisany ręcznie – z pliku "
                    f"{nazwa}_gatunki.csv.",
        "hd": "Smukłość: wysokość / pierśnica (te same jednostki).",
        "d13_uw": "Uwagi z pomiaru pierśnicy w kroku 1 (ok = bez zastrzeżeń).",
        "d13_x": "Pierśnica zmierzona drugi raz przez LidarTreeSegmenter ze śledzenia "
                 "całej strzały – do porównania z D1,3.",
        "d13_x_zr": "pomiar – okrąg dopasowany do punktów na 1,3 m; "
                    "rekonstrukcja – dół pnia zasłonięty, wartość z "
                    "przedłużenia strzały widocznej wyżej.",
        "pien_od": "Od jakiej wysokości skaner widział pień.",
        "pien_do": "Do jakiej wysokości udało się prześledzić pień.",
        "dor_od": "Początek dorysowanego (zasłoniętego) odcinka pnia.",
        "dor_do": "Koniec dorysowanego odcinka pnia.",
        "vz": "Objętość odcinka strzały widzianego przez skaner (suma walców "
              "ściętych). Bez wierzchołka, więc mniejsza od V.",
        "proc_pien": "Udział punktów drzewa zaklasyfikowanych jako pień.",
        "proc_gal": "Udział punktów zaklasyfikowanych jako gałęzie (drewno "
                    "poza pniem).",
        "proc_lis": "Udział punktów zaklasyfikowanych jako ulistnienie / igły.",
        "pkt": "Liczba punktów chmury należących do drzewa.",
        "sc_odl": "Najmniejsza odległość pnia od ścieżki skanera (w "
                  "poziomie). Im dalej, tym mniej punktów na korze i "
                  f"mniej pewna pierśnica (zwłaszcza ponad "
                  f"{SC.DALEKO_M:.0f} m).",
        "sc_obej": "Z jakiej części obwodu (sektory po 10°) ścieżka "
                   "przeszła bliżej niż 15 m. Mało (np. poniżej 30%) = pień "
                   "widziany głównie z jednej strony; na obrazku drzewa "
                   "pokazuje to szary pierścień.",
    }
    naglowek(ws, 3, [t("Kolumna", "Column"), t("Znaczenie", "Meaning")], 22)
    for i, (k, nagl, _, _) in enumerate(kol, start=4):
        ws.cell(row=i, column=1, value=nagl).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=obj.get(k, ""))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        for j in (1, 2):
            ws.cell(row=i, column=j).border = RAMKA
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100

    wyj = rap / f_rap
    try:
        wb.save(wyj)
    except PermissionError:
        wyj = rap / f_rap.replace(".xlsx",
                                  f"_{dt.datetime.now():%H%M%S}.xlsx")
        wb.save(wyj)
        print(t("  (poprzedni raport jest otwarty w Excelu - zapisuję pod "
                "nową nazwą)", "  (the previous report is open in Excel - "
                "saving under a new name)"))
    print(t(f"[RAPORT] Zapisano: {wyj.name} i 00_PODSUMOWANIE.txt",
            f"[REPORT] Saved: {wyj.name} and 00_SUMMARY.txt"), flush=True)
    wszystkie = pd.concat([T, Tk], ignore_index=True).sort_values("nr")
    wszystkie.to_csv(U.katalog(kat, U.TABELE) / f"{nazwa}_drzewa_raport.csv",
                     index=False)
    gpkg_drzewa(kat, nazwa, wszystkie, et, (kr, kr_siatka))
    if teren:
        gpkg_georef(kat, nazwa, teren)
    print(t(f"MIAZSZOSC: {v.sum() / ha:.0f} m3/ha",
            f"VOLUME: {v.sum() / ha:.0f} m3/ha"), flush=True)


def OBJASNIENIA_EN(tablice, wzory, f_gat):
    """Arkusz Explanations w wersji angielskiej."""
    return {
        "nr": "Tree number – the same as in the point cloud files and in the "
              "images (04_OBRAZY_DRZEW\\drzewo_NNN.png).",
        "x": "Stem position in the scan coordinate system.",
        "y": "Stem position in the scan coordinate system.",
        "d13": "Diameter at breast height (1.3 m), measured in step 1 (tree "
               "detection). Basal area G and volume V are based on it.",
        "h": "Tree height – the highest point of the crown above ground.",
        "g": "Basal area at breast height: π/4 · D1.3².",
        "v": f"Merchantable timber volume (stem wood ≥7 cm top diameter, "
             f"over bark) by the formula of the species group in {tablice} "
             f"({wzory}).",
        "z_gory": "yes – stem found only in the second search, higher up "
                  "(base hidden by understory).",
        "grupa": "Species that determines the volume formula. A detected "
                 "coniferous tree is taken as spruce, a broadleaved one as "
                 "beech (the main species on the plots). Enter any other "
                 "species manually.",
        "gr_pew": "How clearly the tree shape indicated the group (rules: "
                  "stem up to the top, sharp apex and conical crown = "
                  "coniferous; forked stem, wide rounded crown = "
                  "broadleaved). Preliminary – check the tree image.",
        "hk": "Height at which the crown begins.",
        "grupa_zr": "default – assumed for all trees; detected – from tree "
                    "shape; scanner label – code saved with the scanner in "
                    "the field (markPoint.csv, e.g. FI = spruce, BU = "
                    f"beech); entered manually – from the file {f_gat}.",
        "hd": "Slenderness: height / DBH (same units).",
        "d13_uw": "Notes on the DBH measurement in step 1 (empty = no "
                  "remarks).",
        "d13_x": f"DBH measured a second time by {J.NAZWA} by tracing the "
                 f"whole stem – for comparison with D1.3.",
        "d13_x_zr": "measured – circle fitted to the points at 1.3 m; "
                    "reconstructed – stem base hidden, value from extending "
                    "the stem visible higher up.",
        "pien_od": "Height from which the scanner saw the stem.",
        "pien_do": "Height up to which the stem could be traced.",
        "dor_od": "Start of the reconstructed (hidden) stem section.",
        "dor_do": "End of the reconstructed stem section.",
        "vz": "Volume of the stem section seen by the scanner (sum of "
              "truncated cones). Without the top, so smaller than V.",
        "proc_pien": "Share of the tree's points classified as stem.",
        "proc_gal": "Share of points classified as branches (wood outside "
                    "the stem).",
        "proc_lis": "Share of points classified as foliage / needles.",
        "pkt": "Number of point cloud points belonging to the tree.",
        "sc_odl": "Smallest horizontal distance of the stem from the scanner "
                  "path. The farther, the fewer points on the bark and the "
                  f"less reliable the DBH (especially beyond "
                  f"{SC.DALEKO_M:.0f} m).",
        "sc_obej": "Share of the circumference (10° sectors) from which the "
                   "path passed closer than 15 m. Low (e.g. under 30%) = "
                   "stem seen mainly from one side; the grey ring on the "
                   "tree image shows this.",
        "sprawdzone": "yes – the tree was checked in the 3D view review "
                      "mode.",
    }


# ----------------------------------------------------------------------
# pomiar terenowy (kilka drzew z GPS) - dopasowanie i georeferencja
# ----------------------------------------------------------------------
def pomiar_terenowy(kat, nazwa, u, T, epsg=None):
    sys.path.insert(0, str(TU))
    import lts_teren as TR
    p = TR.znajdz_pomiar(kat, nazwa, u.get("plik"))
    if p is None:
        return None
    print(t(f"[TEREN] Pomiar terenowy: {p.name}",
            f"[FIELD] Field measurement: {p.name}"), flush=True)
    try:
        teren, uk = TR.wczytaj_pomiar(p, epsg)
        sk = T[~T.krzew.astype(bool)][["nr", "x", "y", "d13"]] \
            .rename(columns={"d13": "d_cm"}).reset_index(drop=True)
        wyn = TR.dopasuj(teren, sk)
    except Exception as e:                                # noqa: BLE001
        print(t(f"  UWAGA: dopasowanie do pomiaru terenowego się nie udało: "
                f"{e}", f"  WARNING: matching to the field measurement "
                f"failed: {e}"), flush=True)
        return None
    ocena = TR.ocena_slowna(wyn, len(teren))
    TR.zapisz_json(U.katalog(kat, U.TABELE) / f"{nazwa}_dopasowanie.json",
                   wyn, uk, teren, sk, ocena)
    print(t(f"[TEREN] Dopasowanie {ocena[0]}: {ocena[1]}; obrót skanu "
            f"{wyn['obrot_deg']:.1f}°, układ EPSG:{uk['epsg']}",
            f"[FIELD] Match {ocena[0]}: {ocena[1]}; scan rotation "
            f"{wyn['obrot_deg']:.1f}°, CRS EPSG:{uk['epsg']}"), flush=True)
    # tabela porównania
    wiersze = []
    odw = {v: k for k, v in wyn["pary"].items()}
    xy_s = TR.przeksztalc(sk[["x", "y"]].to_numpy(float), wyn["R"], wyn["t"])
    for i, r in teren.iterrows():
        j = wyn["pary"].get(i)
        rt = T[T.nr == sk.nr[j]].iloc[0] if j is not None else None
        wiersze.append({
            "nr_teren": int(r.nr),
            "nr_skan": int(sk.nr[j]) if j is not None else None,
            "gat_teren": J.gatunek(r.gatunek), "gat_skan": J.gatunek(rt.grupa)
            if rt is not None else "",
            "d_teren": r.d_cm, "d_skan": rt.d13 if rt is not None else np.nan,
            "h_teren": r.h_m, "h_skan": rt.h if rt is not None else np.nan,
            "odl": float(np.hypot(*(xy_s[j] - [r.x, r.y])))
            if j is not None else np.nan})
    por = pd.DataFrame(wiersze)
    por["d_roz"] = por.d_skan - por.d_teren
    por["h_roz"] = por.h_skan - por.h_teren
    return {"plik": p.name, "teren": teren, "uklad": uk, "wyn": wyn,
            "ocena": ocena, "por": por, "skan": sk, "odw": odw}


def arkusz_krzewow(wb, kr):
    ws = wb.create_sheet(t("Krzewy", "Shrubs"))
    tytul(ws, t("Krzewy (podszyt 0,3–4 m)", "Shrubs (understory 0.3–4 m)"),
          t(f"{len(kr)} krzewów, razem {pl(kr.pole_m2.sum(), 1)} m² · "
            f"wysokość = 95% punktów krzewu niżej",
            f"{len(kr)} shrubs, total {kr.pole_m2.sum():.1f} m² · "
            f"height = 95% of the shrub points lower"))
    kol = [("nr", t("Nr krzewu", "Shrub no."), "0", 9),
           ("x", "X [m]", "0.00", 9), ("y", "Y [m]", "0.00", 9),
           ("pole_m2", t("Pole z góry [m²]", "Area from above [m²]"),
            "0.00", 12),
           ("wysokosc_m", t("Wysokość [m]", "Height [m]"), "0.0", 11),
           ("wysokosc_max_m", t("Najwyższy punkt [m]", "Highest point [m]"),
            "0.0", 12),
           ("punktow", t("Punktów", "Points"), "0", 10)]
    naglowek(ws, 4, [k[1] for k in kol])
    for r, (_, w) in enumerate(kr.sort_values("nr").iterrows(), start=5):
        for c, (k, _, f, _) in enumerate(kol, start=1):
            v = w[k]
            cell = ws.cell(row=r, column=c,
                           value=None if pd.isna(v) else float(v)
                           if k not in ("nr", "punktow") else int(v))
            cell.number_format = f
            cell.border = RAMKA
    for i, k in enumerate(kol):
        ws.column_dimensions[chr(65 + i)].width = k[3]
    ws.freeze_panes = "A5"


def arkusz_terenowy(wb, tr):
    ws = wb.create_sheet(t("Pomiar terenowy", "Field measurement"))
    tytul(ws, t("Porównanie z pomiarem terenowym",
                "Comparison with the field measurement"),
          t(f"plik {tr['plik']} · dopasowanie {tr['ocena'][0]}: "
            f"{tr['ocena'][1]} · obrót skanu "
            f"{pl(tr['wyn']['obrot_deg'], 1)}° · "
            f"układ EPSG:{tr['uklad']['epsg']}",
            f"file {tr['plik']} · match {tr['ocena'][0]}: "
            f"{tr['ocena'][1]} · scan rotation "
            f"{pl(tr['wyn']['obrot_deg'], 1)}° · "
            f"CRS EPSG:{tr['uklad']['epsg']}"))
    kol = [("nr_teren", t("Nr w terenie", "Field no."), "0", 9),
           ("nr_skan", t("Nr w skanie", "Scan no."), "0", 9),
           ("gat_teren", t("Gatunek – teren", "Species – field"), "@", 13),
           ("gat_skan", t("Gatunek – skan", "Species – scan"), "@", 13),
           ("d_teren", t("Pierśnica – teren [cm]", "DBH – field [cm]"),
            "0.0", 11),
           ("d_skan", t("Pierśnica – skan [cm]", "DBH – scan [cm]"), "0.0",
            11),
           ("d_roz", t("Różnica pierśnic skan − teren [cm]", "DBH "
                       "difference scan − field [cm]"), "+0.0;-0.0;0.0", 13),
           ("h_teren", t("Wysokość – teren [m]", "Height – field [m]"),
            "0.0", 11),
           ("h_skan", t("Wysokość – skan [m]", "Height – scan [m]"), "0.0",
            11),
           ("h_roz", t("Różnica wysokości skan − teren [m]", "Height "
                       "difference scan − field [m]"), "+0.0;-0.0;0.0", 13),
           ("odl", t("Odległość od pozycji GPS [m]", "Distance from GPS "
                     "position [m]"), "0.0", 12)]
    naglowek(ws, 4, [k[1] for k in kol], 48)
    por = tr["por"]
    for r, (_, w) in enumerate(por.iterrows(), start=5):
        for c, (k, _, fmt, _) in enumerate(kol, start=1):
            v = w[k]
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            if k == "nr_skan" and v is None:
                v = t("brak pary", "no match")
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = fmt
            cell.border = RAMKA
    for c, (_, _, _, sz) in enumerate(kol, start=1):
        ws.column_dimensions[get_column_letter(c)].width = sz
    r = 6 + len(por)
    dd = por.d_roz.dropna()
    if len(dd):
        ws.cell(row=r, column=1, value=t(
            f"Pierśnica: średnia różnica {pl(dd.mean(), 1)} cm, "
            f"średni błąd bezwzględny {pl(dd.abs().mean(), 1)} cm "
            f"({len(dd)} drzew).",
            f"DBH: mean difference {pl(dd.mean(), 1)} cm, "
            f"mean absolute error {pl(dd.abs().mean(), 1)} cm "
            f"({len(dd)} trees)."))
    ws.cell(row=r + 1, column=1, value=t(
        "Odległość od pozycji GPS to głównie błąd odbiornika pod okapem. "
        "Skan został obrócony i przesunięty tak, żeby te odległości były "
        "jak najmniejsze, a pierśnice zgodne.",
        "The distance from the GPS position is mostly receiver error under "
        "the canopy. The scan was rotated and shifted so that these "
        "distances are as small as possible and the DBH values agree."))
    ws.cell(row=r + 2, column=1, value=t(
        "Wyniki we współrzędnych terenowych są w pliku "
        "<skan>_georef.gpkg (QGIS).", "Results in field coordinates are in "
        "the file <scan>_georef.gpkg (QGIS)."))
    for k in range(r, r + 3):
        ws.cell(row=k, column=1).font = Font(italic=True, color=SZARY)


def gpkg_georef(kat, nazwa, tr):
    """<skan>_georef.gpkg: wszystkie warstwy we współrzędnych terenowych
    plus warstwa pomiar_terenowy z pozycjami GPS i odchyłkami."""
    import lts_gpkg as G
    import lts_teren as TR
    we = U.plik(kat, f"{nazwa}.gpkg")
    if not we.exists():
        return
    wy = we.with_name(f"{nazwa}_georef.gpkg")
    uk = tr["uklad"]
    srs = (uk["epsg"], TR.wkt(uk) if "lon0" in uk else "undefined")
    try:
        G.przeksztalc_plik(we, wy, tr["wyn"]["R"], tr["wyn"]["t"], srs)
        te, por = tr["teren"], tr["por"]
        at = por.copy()
        at["dopasowanie"] = tr["ocena"][0]
        G.zapisz_warstwe(wy, "pomiar_terenowy", "POINT",
                         list(zip(te.x, te.y)), at, srs=srs,
                         opis=t("drzewa zmierzone w terenie (pozycja GPS)",
                                "trees measured in the field (GPS position)"))
        linie, atl = [], []
        xy_s = TR.przeksztalc(tr["skan"][["x", "y"]].to_numpy(float),
                              tr["wyn"]["R"], tr["wyn"]["t"])
        for i, j in tr["wyn"]["pary"].items():
            linie.append(np.array([[te.x[i], te.y[i]], xy_s[j]]))
            atl.append({"nr_teren": int(te.nr[i]),
                        "nr_skan": int(tr["skan"].nr[j]),
                        "odleglosc_m": float(np.hypot(*(xy_s[j] -
                                                        [te.x[i], te.y[i]])))})
        G.zapisz_warstwe(wy, "odchylki_gps", "LINESTRING", linie,
                         pd.DataFrame(atl), srs=srs,
                         opis=t("od pozycji GPS do drzewa w skanie",
                                "from the GPS position to the tree in the "
                                "scan"))
        import lts_qgis as Q
        Q.zapisz_style(wy)
        Q.zapisz_projekt(wy)
        print(t(f"[TEREN] GeoPackage we współrzędnych terenowych: {wy.name}",
                f"[FIELD] GeoPackage in field coordinates: {wy.name}"),
              flush=True)
    except Exception as e:                                # noqa: BLE001
        print(t(f"  UWAGA: nie udało się zapisać {wy.name}: {e}",
                f"  WARNING: could not save {wy.name}: {e}"), flush=True)


def gpkg_drzewa(kat, nazwa, T, et=None, krzewy=None):
    """Podmienia warstwę drzewa w <skan>.gpkg na pełne dane z raportu."""
    plik = U.plik(kat, f"{nazwa}.gpkg")
    if not plik.exists():
        return                       # krok 2 jeszcze nie zrobił GeoPackage
    nazwy = {"nr": "nr", "x": "x", "y": "y", "d13": "piersnica_cm",
             "d13_uw": "uwagi", "h": "wysokosc_m", "hd": "smuklosc_h_d",
             "g": "pole_przekroju_m2", "v": "miazszosc_grubizny_m3",
             "grupa": "gatunek", "grupa_zr": "gatunek_zrodlo",
             "gr_pew": "gatunek_pewnosc_proc", "z_gory": "wykryte_od_gory",
             "d13_x": "piersnica_lts_cm", "d13_x_zr": "piersnica_lts_zrodlo",
             "pien_od": "pien_od_m", "pien_do": "pien_do_m",
             "dor_od": "dorysowano_od_m", "dor_do": "dorysowano_do_m",
             "proc_pien": "proc_pien", "proc_gal": "proc_galezie",
             "proc_lis": "proc_ulistnienie", "pkt": "liczba_punktow",
             "hk": "podstawa_korony_m", "vz": "miazszosc_strzaly_zmierzonej_m3",
             "krzew": "krzew", "kontrola": "kontrola_pnia",
             "sprawdzone": "sprawdzone", "sc_odl": "odl_od_sciezki_m",
             "sc_obej": "obejscie_ze_sciezki_proc"}
    A = T[[k for k in nazwy if k in T]].rename(columns=nazwy).copy()
    A["krzew"] = A.krzew.fillna(False).astype(bool)
    try:
        sys.path.insert(0, str(TU))
        import lts_gpkg as G
        srs = G.wczytaj_srs(U.plik(kat, f"{nazwa}_uklad.json"))
        G.napraw_uklad(plik)         # starsze pliki: układ -1 -> metry
        # starsze pliki: warstwy, których już nie zapisujemy
        us = G.usun_warstwy(plik, G.ZBEDNE)
        if us:
            print(t("         usunięte zbędne warstwy: ",
                    "         removed unneeded layers: ") + ", ".join(us),
                  flush=True)
        # pliki ze starszych wersji: do 3000 punktów plastra na drzewo
        for nw in ("punkty_1_3m", "punkty_3m", "points_1_3m", "points_3m"):
            pp = G.przerzedz_warstwe_punktow(plik, nw, srs)
            if pp:
                print(t(f"         {nw}: przerzedzone z {pp[0]:,} do "
                        f"{pp[1]:,} punktów".replace(",", " "),
                        f"         {nw}: thinned from {pp[0]:,} to "
                        f"{pp[1]:,} points"), flush=True)
        G.zapisz_warstwe(plik, "drzewa", "POINT",
                         list(zip(T.x.astype(float), T.y.astype(float))), A,
                         srs=srs, opis=t("drzewa: dane z raportu LidarTreeSegmenter",
                                         f"trees: data from the {J.NAZWA} "
                                         f"report"))
        sc = SC.dla_skanu(kat, nazwa, cicho=True,
                          xy_skanu=T[["x", "y"]].to_numpy(float))
        if sc:
            xyz = SC.uprosc(sc["xyz"], 0.25)
            G.zapisz_warstwe(plik, "sciezka_skanera", "LINESTRING",
                             [[tuple(p) for p in xyz]],
                             pd.DataFrame({"plik": [sc["plik"].name],
                                           "dlugosc_m": [round(
                                               sc["dlugosc_m"], 1)],
                                           "punktow": [len(sc["xyz"])]}),
                             z=True, srs=srs,
                             opis=t("ścieżka skanera (trajektoria)",
                                    "scanner path (trajectory)"))
        if et and et["wszystkie"]:
            E = et["wszystkie"]
            G.zapisz_warstwe(plik, "etykiety_skanera", "POINT",
                             [(e["x"], e["y"]) for e in E],
                             pd.DataFrame({
                                 "kod": [e["kod"] for e in E],
                                 "gatunek": [J.gatunek(e["gatunek"] or "")
                                             for e in E],
                                 "nr_drzewa": [e["nr"] for e in E],
                                 "odleglosc_m": [round(e["odl"], 2)
                                                 for e in E],
                                 "przypisana": [J.TAK if e["nr"] is not None
                                                and e["gatunek"] else ""
                                                for e in E],
                                 "czas": [e["czas"] for e in E]}),
                             srs=srs,
                             opis=t("etykiety gatunku zapisane skanerem",
                                    "species labels saved by the scanner"))
        kr, ks = krzewy or (None, None)
        if kr is not None and len(kr) and ks is not None:
            geo, wier = [], []
            for _, w in kr.iterrows():
                o = KR.obrys(ks, int(w.nr))
                if o is not None and len(o) >= 3:
                    geo.append(o)
                    wier.append(w)
            if geo:
                A_k = pd.DataFrame(wier).reset_index(drop=True)
                A_k.columns = ["nr", "x", "y", "pole_m2", "wysokosc_m",
                               "wysokosc_max_m", "punktow"]
                G.zapisz_warstwe(plik, "krzewy", "POLYGON", geo, A_k,
                                 srs=srs, opis=t("krzewy (podszyt 0,3-4 m)",
                                                 "shrubs (understory "
                                                 "0.3-4 m)"))
        import lts_qgis as Q
        Q.zapisz_style(plik)
        qgs = Q.zapisz_projekt(plik)
        print(t(f"[RAPORT] GeoPackage: warstwa drzewa uzupełniona "
                f"({plik.name}), style warstw i projekt QGIS "
                f"{qgs.name if qgs else ''}",
                f"[REPORT] GeoPackage: layer trees updated ({plik.name}), "
                f"layer styles and QGIS project "
                f"{qgs.name if qgs else ''}"), flush=True)
    except Exception as e:                                # noqa: BLE001
        print(t(f"  UWAGA: nie udało się uzupełnić {plik.name}: {e}",
                f"  WARNING: could not update {plik.name}: {e}"),
              flush=True)


if __name__ == "__main__":
    main()
